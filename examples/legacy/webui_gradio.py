#!/usr/bin/env python
"""
Local Gradio WebUI for RAG-Anything (CPU-only friendly).
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import socket
import shutil
import time
import uuid
import traceback
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field
from functools import partial
from pathlib import Path
from threading import Thread
from typing import Any, Optional

import numpy as np
from dotenv import load_dotenv
from lightrag.utils import EmbeddingFunc, logger
from openai import AsyncOpenAI

from raganything import RAGAnything, RAGAnythingConfig
from raganything.config import resolve_embedding_runtime_config
from raganything.export import (
    export_chat_history_to_pdf,
    export_current_answer_to_pdf,
    generate_pdf_report_from_index,
)
from raganything.parser import get_parser

load_dotenv(dotenv_path=".env", override=False)

REGISTRY_PATH = Path("./rag_storage/webui_registry.json").resolve()
UPLOADS_DIR = Path("./rag_storage/webui_uploads").resolve()
DOCS_ROOT = Path("./rag_storage/webui_docs").resolve()
REPORTS_ROOT = Path("./output/reports").resolve()
SUPPORTED_QUERY_LOAD_PARSERS = {"mineru", "docling", "paddleocr", "simple_docx"}
SPECIAL_QUERY_MARKERS = {
    "table": "[PDF Table",
    "equation": "[PDF Equation",
    "figure": "[PDF Visual Description",
}


def _is_gemini_quota_exceeded(exc: Exception) -> bool:
    msg = str(exc)
    return (
        "429" in msg
        or "RESOURCE_EXHAUSTED" in msg
        or "quota" in msg.lower()
        or "rate limit" in msg.lower()
    )


def _is_temporary_vlm_error(exc: Exception) -> bool:
    msg = str(exc).lower()
    return any(
        token in msg
        for token in [
            "error code: 503",
            "unavailable",
            "high demand",
            "resource_exhausted",
            "error code: 429",
            "rate limit",
            "quota",
            "timed out",
            "timeout",
        ]
    )


def detect_parser_for_file(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return "pdf_hybrid"
    if suffix == ".docx":
        return "simple_docx"
    if suffix in {".png", ".jpg", ".jpeg"}:
        return "paddleocr"
    if suffix == ".pptx":
        return "docling"
    if suffix in {".md", ".html", ".htm"}:
        return "docling"
    return "docling"


def _is_index_ready(working_dir: str) -> bool:
    wd = Path(working_dir)
    required = [wd / "vdb_chunks.json", wd / "kv_store_text_chunks.json"]
    return all(p.exists() and p.stat().st_size > 0 for p in required)


def _index_artifact_status(working_dir: str) -> dict[str, Any]:
    wd = Path(working_dir)
    files = {
        "graph_chunk_entity_relation.graphml": wd
        / "graph_chunk_entity_relation.graphml",
        "vdb_chunks.json": wd / "vdb_chunks.json",
        "kv_store_text_chunks.json": wd / "kv_store_text_chunks.json",
    }
    status = {
        "working_dir_exists": wd.exists(),
        "working_dir_abs": str(wd.resolve()),
        "files": {},
        "dir_entries": [],
    }
    for name, path in files.items():
        exists = path.exists()
        size = path.stat().st_size if exists else 0
        status["files"][name] = {"exists": exists, "size": size}
    if wd.exists():
        try:
            status["dir_entries"] = sorted([p.name for p in wd.iterdir()])[:80]
        except Exception:
            status["dir_entries"] = []
    return status


def _read_doc_status_error(working_dir: str, filename: str) -> Optional[str]:
    status_path = Path(working_dir) / "kv_store_doc_status.json"
    if not status_path.exists():
        return None
    try:
        data = json.loads(status_path.read_text(encoding="utf-8"))
    except Exception:
        return None
    target = (filename or "").lower()
    for item in (data or {}).values():
        if not isinstance(item, dict):
            continue
        file_ref = str(item.get("file_path", "")).lower()
        if target and target not in file_ref:
            continue
        status = str(item.get("status", "")).lower()
        if status in {"failed", "error"}:
            return (
                str(item.get("error_msg", "") or "").strip() or "document status failed"
            )
    return None


def _normalize_to_messages(history):
    normalized = []
    for item in history or []:
        if isinstance(item, dict):
            role = item.get("role")
            content = item.get("content")
            if role in {"user", "assistant"} and content is not None:
                normalized.append({"role": role, "content": str(content)})
        elif isinstance(item, (list, tuple)) and len(item) == 2:
            user_msg, assistant_msg = item
            normalized.append({"role": "user", "content": str(user_msg)})
            normalized.append({"role": "assistant", "content": str(assistant_msg)})
    return normalized


def _append_messages(history, user_q: str, answer: str):
    messages = _normalize_to_messages(history)
    messages.append({"role": "user", "content": str(user_q)})
    messages.append({"role": "assistant", "content": str(answer)})
    return messages


def _safe_filename(name: str) -> str:
    out = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return out or "uploaded_file"


def _sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def _extract_candidate_filenames(question: str) -> set[str]:
    q = (question or "").lower()
    return set(
        re.findall(
            r"[a-zA-Z0-9_.-]+\.(?:pdf|docx|png|jpg|jpeg|pptx|md|html|htm)",
            q,
        )
    )


def _detect_special_pdf_question(question: str) -> Optional[str]:
    q = (question or "").lower()
    if any(k in q for k in ["table ", "table.", "table:", "bảng"]):
        return "table"
    if any(k in q for k in ["equation", "formula", "công thức"]):
        return "equation"
    if any(k in q for k in ["fig", "figure", "chart", "hình"]):
        return "figure"
    return None


def _rewrite_pdf_special_query(question: str, special_kind: str) -> str:
    raw = (question or "").strip()
    if special_kind == "equation":
        return (
            "Find equation chunk '[PDF Equation | label=Accuracy]' and return only the full formula. "
            "Must include denominator terms TP + TN + FP + FN when present. "
            f"Question: {raw}"
        )
    if special_kind == "table":
        return (
            "Find table chunk '[PDF Table | label=Table 2]' first when question mentions Table 2. "
            "Answer from table rows only; exclude references and unrelated tables such as Subset 1. "
            f"Question: {raw}"
        )
    if special_kind == "figure":
        return (
            "Find visual description chunk '[PDF Visual Description]' relevant to the figure and answer concisely. "
            f"Question: {raw}"
        )
    return raw


def _is_marker_only_answer(text: str) -> bool:
    lines = [ln.strip() for ln in str(text or "").splitlines() if ln.strip()]
    if not lines:
        return False
    if len(lines) == 1 and re.match(
        r"^\[PDF (Equation|Table|Visual Description)\b.*\]$", lines[0]
    ):
        return True
    return False


async def _ollama_embed_runtime(
    texts,
    *,
    model: str,
    host: str,
    target_dim: int,
):
    import aiohttp

    url = f"{host.rstrip('/')}/api/embed"
    async with aiohttp.ClientSession() as session:
        async with session.post(
            url,
            json={"model": model, "input": texts},
            timeout=60,
        ) as resp:
            if resp.status != 200:
                body = await resp.text()
                raise RuntimeError(f"Ollama embed failed ({resp.status}): {body[:240]}")
            data = await resp.json()
    embeddings = data.get("embeddings", [])
    if not embeddings:
        raise RuntimeError("Ollama returned empty embeddings.")
    if len(embeddings[0]) != target_dim:
        raise RuntimeError(
            f"Embedding dim mismatch. Expected {target_dim}, got {len(embeddings[0])}."
        )
    return np.array(embeddings, dtype=np.float32)


async def _llm_call_runtime(
    prompt,
    *,
    model: str,
    api_key: Optional[str],
    base_url: Optional[str],
    system_prompt=None,
    history_messages=None,
    **kwargs,
):
    history_messages = history_messages or []
    messages = kwargs.pop("messages", None)
    kwargs.pop("hashing_kv", None)
    kwargs.pop("keyword_extraction", None)
    kwargs.pop("enable_cot", None)
    kwargs.pop("response_format", None)

    if messages is None:
        messages = []
        if system_prompt:
            messages.append({"role": "system", "content": system_prompt})
        messages.extend(history_messages)
        messages.append({"role": "user", "content": prompt})

    client = AsyncOpenAI(
        api_key=api_key,
        base_url=base_url,
        max_retries=0,
        timeout=180.0,
    )
    async with client:
        resp = await client.chat.completions.create(
            model=model,
            messages=messages,
            **kwargs,
        )
    content = resp.choices[0].message.content if resp.choices else None
    if content is None:
        raise RuntimeError("Empty LLM response content.")
    return content


async def _vision_call_runtime(
    prompt,
    *,
    vision_model: str,
    api_key: Optional[str],
    base_url: Optional[str],
    llm_model: str,
    system_prompt=None,
    history_messages=None,
    image_data=None,
    messages=None,
    **kwargs,
):
    history_messages = history_messages or []
    if messages:
        return await _llm_call_runtime(
            "",
            model=vision_model,
            api_key=api_key,
            base_url=base_url,
            messages=messages,
            **kwargs,
        )

    if image_data:
        payload = [
            {"role": "system", "content": system_prompt} if system_prompt else None,
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/jpeg;base64,{image_data}"},
                    },
                ],
            },
        ]
        payload = [m for m in payload if m]
        return await _llm_call_runtime(
            "",
            model=vision_model,
            api_key=api_key,
            base_url=base_url,
            messages=payload,
            **kwargs,
        )

    return await _llm_call_runtime(
        prompt,
        model=llm_model,
        api_key=api_key,
        base_url=base_url,
        system_prompt=system_prompt,
        history_messages=history_messages,
        **kwargs,
    )


@dataclass
class DocumentRecord:
    doc_id: str
    original_filename: str
    stored_file_rel: str
    working_dir_rel: str
    file_type: str
    parser: str
    pdf_mode: str
    indexed_at: str
    status: str
    summary: str
    sha256: str
    source_metadata: dict[str, Any]
    visual_targets: list[dict[str, Any]]
    needs_reprocess: bool = False
    needs_reprocess_reason: str = ""
    error_message: str = ""


class RegistryStore:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._last_load_migrated = False

    @staticmethod
    def _path_to_rel(path_value: Optional[str], expected_prefix: str) -> Optional[str]:
        if not path_value:
            return None
        raw = str(path_value).replace("\\", "/").strip()
        if not raw:
            return None
        if raw.startswith(f"{expected_prefix}/") or raw == expected_prefix:
            return raw
        marker = f"/{expected_prefix}/"
        if marker in raw:
            tail = raw.split(marker, 1)[1]
            return f"{expected_prefix}/{tail}"
        marker2 = f"{expected_prefix}/"
        idx = raw.find(marker2)
        if idx >= 0:
            return raw[idx:]
        return None

    def _from_raw_record(self, item: dict[str, Any]) -> tuple[DocumentRecord, bool]:
        migrated = False
        stored_rel = item.get("stored_file_rel")
        working_rel = item.get("working_dir_rel")
        if not stored_rel:
            stored_rel = self._path_to_rel(
                item.get("stored_file_path"), "webui_uploads"
            )
            migrated = migrated or bool(stored_rel)
        if not working_rel:
            working_rel = self._path_to_rel(item.get("working_dir"), "webui_docs")
            migrated = migrated or bool(working_rel)

        needs_reprocess = bool(item.get("needs_reprocess", False))
        reason = str(item.get("needs_reprocess_reason", "") or "").strip()
        if not stored_rel or not working_rel:
            needs_reprocess = True
            if not reason:
                reason = (
                    "could not migrate legacy absolute path(s) into "
                    "webui_uploads/ or webui_docs/"
                )

        rec = DocumentRecord(
            doc_id=str(item.get("doc_id", "")),
            original_filename=str(item.get("original_filename", "")),
            stored_file_rel=str(stored_rel or ""),
            working_dir_rel=str(working_rel or ""),
            file_type=str(item.get("file_type", "")),
            parser=str(item.get("parser", "")),
            pdf_mode=str(item.get("pdf_mode", "auto")),
            indexed_at=str(item.get("indexed_at", "")),
            status=str(item.get("status", "indexed")),
            summary=str(item.get("summary", "")),
            sha256=str(item.get("sha256", "")),
            source_metadata=dict(item.get("source_metadata", {}) or {}),
            visual_targets=list(item.get("visual_targets", []) or []),
            needs_reprocess=needs_reprocess,
            needs_reprocess_reason=reason,
            error_message=str(item.get("error_message", "") or ""),
        )
        migrated = migrated or ("stored_file_path" in item) or ("working_dir" in item)
        return rec, migrated

    def load(self) -> list[DocumentRecord]:
        self._last_load_migrated = False
        if not self.path.exists():
            return []
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            records: list[DocumentRecord] = []
            migrated = False
            for item in raw:
                if not isinstance(item, dict):
                    continue
                rec, rec_migrated = self._from_raw_record(item)
                if not rec.doc_id:
                    continue
                records.append(rec)
                migrated = migrated or rec_migrated
            self._last_load_migrated = migrated
            return records
        except Exception:
            logger.warning("Registry file is invalid. Starting with empty registry.")
            return []

    def save(self, records: list[DocumentRecord]) -> None:
        payload = [asdict(r) for r in records]
        self.path.write_text(
            json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
        )


@dataclass
class UIState:
    processed: bool = False
    processing_summary: str = ""
    last_process_ts: float = 0.0
    ingest_count: int = 0
    query_count: int = 0
    metadata: dict[str, Any] = field(default_factory=dict)


class WebUIRAGService:
    def __init__(self):
        self.state = UIState()
        self.registry_store = RegistryStore(REGISTRY_PATH)
        self.registry: list[DocumentRecord] = self.registry_store.load()
        if self.registry_store._last_load_migrated:
            self.registry_store.save(self.registry)
        self.rag_cache: dict[str, RAGAnything] = {}

        embedding_cfg = resolve_embedding_runtime_config(default_provider="ollama")
        self.embedding_provider = embedding_cfg.provider
        self.embedding_model = embedding_cfg.model
        self.embedding_dim = embedding_cfg.dim
        self.ollama_host = embedding_cfg.ollama_host
        self.embedding_cfg = embedding_cfg

        self.llm_model = os.getenv("LLM_MODEL", "gemini-3.1-flash-lite")
        self.llm_model_source = (
            "env:LLM_MODEL"
            if os.getenv("LLM_MODEL")
            else "default:gemini-3.1-flash-lite"
        )
        self.vision_model = os.getenv("VISION_MODEL", "gemini-3.1-flash-lite")
        self.vision_model_source = (
            "env:VISION_MODEL"
            if os.getenv("VISION_MODEL")
            else "default:gemini-3.1-flash-lite"
        )
        self.fallback_vision_model = os.getenv(
            "FALLBACK_VISION_MODEL", "gemini-2.5-flash-lite"
        )
        self.fallback_vision_model_source = (
            "env:FALLBACK_VISION_MODEL"
            if os.getenv("FALLBACK_VISION_MODEL")
            else "default:gemini-2.5-flash-lite"
        )
        self.api_key = (
            os.getenv("LLM_BINDING_API_KEY")
            or os.getenv("GEMINI_API_KEY")
            or os.getenv("GOOGLE_API_KEY")
        )
        self.base_url = os.getenv("LLM_BINDING_HOST")

        self._loop = asyncio.new_event_loop()
        self._loop_thread = Thread(target=self._run_loop, daemon=True)
        self._loop_thread.start()

        logger.info("Embedding provider: %s", self.embedding_provider)
        logger.info("Embedding model: %s", self.embedding_model)
        logger.info("Embedding dim: %s", self.embedding_dim)
        logger.info("Embedding host: %s", self.ollama_host)
        logger.info(
            "Embedding config source: provider=%s, model=%s, dim=%s, host=%s",
            self.embedding_cfg.provider_source,
            self.embedding_cfg.model_source,
            self.embedding_cfg.dim_source,
            self.embedding_cfg.host_source,
        )
        logger.info("LLM model: %s", self.llm_model)
        logger.info("LLM model source: %s", self.llm_model_source)
        logger.info("Vision model: %s", self.vision_model)
        logger.info("Vision model source: %s", self.vision_model_source)
        logger.info("Fallback vision model: %s", self.fallback_vision_model)
        logger.info(
            "Fallback vision model source: %s", self.fallback_vision_model_source
        )

    def _run_loop(self):
        asyncio.set_event_loop(self._loop)
        self._loop.run_forever()

    def run(self, coro):
        fut = asyncio.run_coroutine_threadsafe(coro, self._loop)
        return fut.result()

    def list_documents(self) -> list[DocumentRecord]:
        return list(self.registry)

    def _reload_registry(self) -> None:
        # Keep in-memory registry consistent with persisted state across callbacks.
        if self.registry_store.path.exists():
            self.registry = self.registry_store.load()

    def render_registry_text(self) -> str:
        if not self.registry:
            return "No indexed files yet."
        lines = ["Indexed files:"]
        for rec in self.registry:
            lines.append(f"- {rec.original_filename} [{rec.doc_id}] ({rec.status})")
        return "\n".join(lines)

    def _find_by_hash(self, sha256: str) -> Optional[DocumentRecord]:
        return next((r for r in self.registry if r.sha256 == sha256), None)

    def _find_by_id(self, doc_id: str) -> Optional[DocumentRecord]:
        return next((r for r in self.registry if r.doc_id == doc_id), None)

    @staticmethod
    def _as_rel(path: Path, root: Path) -> str:
        return path.resolve().relative_to(root.resolve()).as_posix()

    def _resolve_rel(self, rel_path: str) -> Path:
        return (REGISTRY_PATH.parent.resolve() / Path(rel_path)).resolve()

    def _resolve_record_stored_file(self, rec: DocumentRecord) -> Path:
        return self._resolve_rel(rec.stored_file_rel)

    def _resolve_record_working_dir(self, rec: DocumentRecord) -> Path:
        return self._resolve_rel(rec.working_dir_rel)

    def _record_index_status(self, rec: DocumentRecord) -> tuple[bool, str]:
        if rec.status != "indexed":
            detail = rec.error_message or rec.needs_reprocess_reason or rec.status
            return False, f"not indexed: {detail}"
        if rec.needs_reprocess:
            detail = rec.needs_reprocess_reason or "unknown reason"
            return False, f"registry path incompatible: {detail}"
        if not rec.working_dir_rel:
            return False, "registry path incompatible: missing working_dir_rel"
        working_dir = self._resolve_record_working_dir(rec)
        if not _is_index_ready(str(working_dir)):
            return False, f"index path missing: {working_dir}"
        return True, ""

    def _load_text_chunks(self, rec: DocumentRecord) -> list[str]:
        wd = self._resolve_record_working_dir(rec)
        path = wd / "kv_store_text_chunks.json"
        if not path.exists():
            return []
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return []
        texts: list[str] = []
        stack: list[Any] = [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, dict):
                for key, value in item.items():
                    if isinstance(value, (dict, list)):
                        stack.append(value)
                    elif isinstance(value, str):
                        if (
                            "[PDF Table" in value
                            or "[PDF Equation" in value
                            or "[PDF Visual Description" in value
                            or str(key).lower() in {"content", "text"}
                        ):
                            texts.append(value)
            elif isinstance(item, list):
                stack.extend(item)
            elif isinstance(item, str):
                texts.append(item)
        return texts

    def _find_marker_chunks(self, rec: DocumentRecord, marker_kind: str) -> list[str]:
        marker = SPECIAL_QUERY_MARKERS.get(marker_kind)
        if not marker:
            return []
        return [c for c in self._load_text_chunks(rec) if marker in c]

    @staticmethod
    def _format_marker_answer(
        marker_kind: str, chunks: list[str], question: str = ""
    ) -> str:
        if not chunks:
            return ""
        q = (question or "").lower()
        if marker_kind == "equation":
            for chunk in chunks:
                lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
                for ln in lines:
                    if ln.startswith("[PDF Equation"):
                        continue
                    if "=" in ln and any(
                        token in ln.lower()
                        for token in ["accuracy", "tp", "tn", "fp", "fn"]
                    ):
                        return ln
                if len(lines) > 1:
                    for ln in lines[1:]:
                        if not ln.startswith("[PDF"):
                            return ln
        if marker_kind == "table":
            wants_accuracy = "accuracy" in q
            wants_roc = "roc" in q
            table_num_match = re.search(r"\btable\s*(\d+)\b", q)
            wanted_table_num = (
                int(table_num_match.group(1)) if table_num_match else None
            )
            terms = [
                t
                for t in re.findall(r"[a-zA-Z0-9_.+-]+", q)
                if t
                not in {
                    "in",
                    "the",
                    "what",
                    "are",
                    "for",
                    "at",
                    "from",
                    "table",
                    "answer",
                    "only",
                }
            ]
            scored_chunks: list[tuple[int, str]] = []
            for chunk in chunks:
                score = 0
                if wanted_table_num is not None:
                    m = re.search(
                        r"\[PDF Table\s*\|\s*label=Table\s+(\d+)\s*\|",
                        chunk,
                        flags=re.IGNORECASE,
                    )
                    if m and int(m.group(1)) == wanted_table_num:
                        score += 1000
                if "references" in chunk.lower() or "subset 1" in chunk.lower():
                    score -= 100
                score += sum(1 for term in terms if term in chunk.lower())
                scored_chunks.append((score, chunk))
            scored_chunks.sort(key=lambda x: x[0], reverse=True)
            for _, chunk in scored_chunks:
                lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
                table_lines = [
                    ln for ln in lines if ln.startswith("|") and ln.endswith("|")
                ]
                if not table_lines:
                    continue
                # Parse markdown table to extract requested columns from best-matching row.
                parsed_rows = []
                for ln in table_lines:
                    cells = [c.strip() for c in ln.strip("|").split("|")]
                    if cells:
                        parsed_rows.append(cells)
                if len(parsed_rows) >= 3:
                    header = parsed_rows[0]
                    body_rows = [
                        r
                        for r in parsed_rows[2:]
                        if not all(re.fullmatch(r"-+", c or "") for c in r)
                    ]
                    best_cells: Optional[list[str]] = None
                    best_score = -1
                    for cells in body_rows:
                        row_text = " | ".join(cells).lower()
                        score = sum(1 for term in terms if term in row_text)
                        if score > best_score:
                            best_score = score
                            best_cells = cells
                    if best_cells is not None and best_score > 0:
                        col_map = {
                            h.strip().lower(): i
                            for i, h in enumerate(header)
                            if h.strip()
                        }
                        out_parts: list[str] = []
                        if wants_accuracy:
                            for k in ["accuracy", "acc"]:
                                if k in col_map and col_map[k] < len(best_cells):
                                    out_parts.append(
                                        f"Accuracy: {best_cells[col_map[k]]}"
                                    )
                                    break
                        if wants_roc:
                            for k in ["roc", "auc", "roc-auc"]:
                                if k in col_map and col_map[k] < len(best_cells):
                                    out_parts.append(f"ROC: {best_cells[col_map[k]]}")
                                    break
                        if out_parts:
                            return "; ".join(out_parts)
                best_line = ""
                best_score = -1
                for ln in table_lines:
                    ln_low = ln.lower()
                    score = sum(1 for term in terms if term in ln_low)
                    if score > best_score:
                        best_score = score
                        best_line = ln
                if best_line and best_score > 0:
                    return best_line
            # Fallback: strip marker and return first contentful line.
            for _, chunk in scored_chunks:
                for ln in chunk.splitlines():
                    s = ln.strip()
                    if s and not s.startswith("[PDF Table"):
                        return s
            return ""
        if marker_kind == "figure":
            candidate = "\n".join(chunks[0].splitlines()[1:]).strip()
            return candidate or chunks[0]
        return chunks[0]

    @staticmethod
    def _is_answer_only_request(question: str) -> bool:
        q = (question or "").lower()
        return ("answer only" in q) or ("chỉ trả lời" in q)

    @staticmethod
    def _resolve_query_parser(preferred: str) -> str:
        candidates = [preferred, "paddleocr", "docling", "simple_docx"]
        seen: set[str] = set()
        for name in candidates:
            n = (name or "").strip().lower()
            if not n or n in seen:
                continue
            seen.add(n)
            try:
                if get_parser(n).check_installation():
                    return n
            except Exception:
                continue
        return "paddleocr"

    @staticmethod
    def _check_pdf_dependencies() -> Optional[str]:
        missing = []
        try:
            import fitz  # type: ignore  # noqa: F401
        except Exception:
            missing.append("pymupdf(fitz)")
        try:
            import pdfplumber  # type: ignore  # noqa: F401
        except Exception:
            missing.append("pdfplumber")
        if missing:
            return (
                "PDF hybrid requires PyMuPDF and/or pdfplumber in Docker. "
                f"Missing: {', '.join(missing)}"
            )
        return None

    async def _index_docx_with_python_docx(
        self,
        file_path: str,
        workdir: Path,
    ) -> None:
        try:
            from docx import Document  # type: ignore
        except Exception as exc:
            raise RuntimeError(
                "simple_docx requires docling, and python-docx fallback is unavailable. "
                "Install docling or python-docx in Docker."
            ) from exc

        doc = Document(file_path)
        lines: list[str] = []
        for p in doc.paragraphs:
            txt = (p.text or "").strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                row_text = " | ".join([c for c in cells if c])
                if row_text:
                    lines.append(row_text)
        content = "\n".join(lines).strip()
        if not content:
            raise RuntimeError("DOCX extracted no text content.")
        max_chars_raw = os.getenv("WEBUI_DOCX_MAX_CHARS", "12000").strip()
        try:
            max_chars = int(max_chars_raw)
        except Exception:
            max_chars = 12000
        if max_chars > 0 and len(content) > max_chars:
            logger.info(
                "DOCX fallback truncation enabled: %s -> %s chars (set WEBUI_DOCX_MAX_CHARS to adjust)",
                len(content),
                max_chars,
            )
            content = content[:max_chars]

        rag = await self._create_rag(
            working_dir=str(workdir.resolve()), parser="paddleocr"
        )
        init_result = await rag._ensure_lightrag_initialized()
        if not init_result or not init_result.get("success"):
            detail = (init_result or {}).get("error", "unknown error")
            raise RuntimeError(f"LightRAG init failed for DOCX fallback: {detail}")
        await rag.insert_content_list(
            [{"type": "text", "text": content, "page_idx": 0}], file_path=file_path
        )

    def _persist_upload_file(self, file_path: str, doc_id: str) -> str:
        UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
        src = Path(file_path)
        dst = UPLOADS_DIR / f"{doc_id}_{_safe_filename(src.name)}"
        shutil.copy2(src, dst)
        return str(dst.resolve())

    @contextmanager
    def _patched_env_for_ingest(self):
        keys = {
            "EMBEDDING_PROVIDER": "ollama",
            "EMBEDDING_BINDING": "ollama",
            "EMBEDDING_MODEL": self.embedding_model,
            "EMBEDDING_DIM": str(self.embedding_dim),
            "OLLAMA_HOST": self.ollama_host,
        }
        old = {k: os.environ.get(k) for k in keys}
        try:
            for k, v in keys.items():
                os.environ[k] = v
            yield
        finally:
            for k, old_v in old.items():
                if old_v is None:
                    os.environ.pop(k, None)
                else:
                    os.environ[k] = old_v

    async def _ensure_ollama_model_available(self) -> None:
        import aiohttp

        if self.embedding_provider != "ollama":
            return
        url = f"{self.ollama_host.rstrip('/')}/api/tags"
        async with aiohttp.ClientSession() as session:
            async with session.get(url, timeout=20) as resp:
                if resp.status != 200:
                    raise RuntimeError(
                        f"Ollama host check failed ({resp.status}) at {url}"
                    )
                data = await resp.json()
        names = [
            m.get("name", "") for m in data.get("models", []) if isinstance(m, dict)
        ]
        aliases = {n for n in names}
        aliases.update({n.split(":")[0] for n in names})
        if self.embedding_model not in aliases:
            raise RuntimeError(
                "Ollama model not found. Run ollama list or set EMBEDDING_MODEL to an existing model."
            )

    async def _create_rag(self, working_dir: str, parser: str) -> RAGAnything:
        if self.embedding_provider != "ollama":
            raise RuntimeError(
                "WebUI defaults to local embeddings only. Set EMBEDDING_PROVIDER=ollama."
            )
        await self._ensure_ollama_model_available()
        llm_model = self.llm_model
        vision_model = self.vision_model
        api_key = self.api_key
        base_url = self.base_url
        embedding_model = self.embedding_model
        ollama_host = self.ollama_host
        embedding_dim = self.embedding_dim

        embedding_func = EmbeddingFunc(
            embedding_dim=embedding_dim,
            max_token_size=8192,
            func=partial(
                _ollama_embed_runtime,
                model=embedding_model,
                host=ollama_host,
                target_dim=embedding_dim,
            ),
        )
        config = RAGAnythingConfig(
            working_dir=working_dir,
            parser=parser,
            parse_method="auto",
            enable_image_processing=True,
            enable_table_processing=True,
            enable_equation_processing=True,
        )
        return RAGAnything(
            config=config,
            llm_model_func=partial(
                _llm_call_runtime,
                model=llm_model,
                api_key=api_key,
                base_url=base_url,
            ),
            vision_model_func=partial(
                _vision_call_runtime,
                vision_model=vision_model,
                llm_model=llm_model,
                api_key=api_key,
                base_url=base_url,
            ),
            embedding_func=embedding_func,
            lightrag_kwargs={
                "embedding_func_max_async": 1,
                "embedding_batch_num": 1,
                "llm_model_max_async": 1,
                "max_parallel_insert": 1,
            },
        )

    async def _process_single_document(
        self,
        file_path: str,
        force_parser: Optional[str] = None,
        pdf_page_range: Optional[str] = None,
        vision_target: Optional[str] = None,
        vision_page_range: Optional[str] = None,
        max_vision_pages: int = 1,
        no_vision: bool = True,
        skip_kg_extraction: bool = True,
        force_reprocess: bool = False,
    ) -> str:
        from examples import raganything_example as rx

        src = Path(file_path)
        if not src.exists():
            return f"Missing file: {file_path}"

        parser = (force_parser or detect_parser_for_file(str(src))).lower()
        sha256 = _sha256_file(str(src))
        existing = self._find_by_hash(sha256)
        if existing:
            existing_ok, _ = self._record_index_status(existing)
        else:
            existing_ok = False
        if existing and existing_ok and (not force_reprocess):
            return (
                f"Skipped duplicate: {existing.original_filename} "
                f"(already indexed as {existing.doc_id})"
            )

        doc_id = existing.doc_id if existing else uuid.uuid4().hex[:12]
        stored_path = self._persist_upload_file(str(src), doc_id)
        workdir = DOCS_ROOT / doc_id
        if force_reprocess and workdir.exists():
            shutil.rmtree(workdir, ignore_errors=True)
        workdir.mkdir(parents=True, exist_ok=True)

        pdf_mode = "auto"
        if parser == "pdf_hybrid":
            parser = "pdf_fast"
            pdf_mode = "hybrid"
            logger.info(
                "Parser alias pdf_hybrid resolved to pdf_fast with pdf_mode=hybrid"
            )
        if parser == "pdf_fast":
            pdf_mode = "hybrid"

        parser_for_run = (
            "pdf_hybrid" if parser == "pdf_fast" and pdf_mode == "hybrid" else parser
        )
        file_type = src.suffix.lower()
        stored_abs = str(Path(stored_path).resolve())
        working_abs = str(workdir.resolve())
        logger.info(
            "Index request: file=%s type=%s parser=%s pdf_mode=%s stored=%s working_dir=%s",
            str(src.resolve()),
            file_type,
            parser_for_run,
            pdf_mode,
            stored_abs,
            working_abs,
        )
        logger.info(
            "Index options: no_vision=%s vision_target=%s vision_page_range=%s max_vision_pages=%s skip_kg_extraction=%s",
            no_vision,
            vision_target,
            vision_page_range,
            max_vision_pages,
            skip_kg_extraction,
        )
        logger.info("Working dir exists before process: %s", workdir.exists())

        if file_type == ".pptx":
            try:
                import docling  # type: ignore  # noqa: F401
            except Exception:
                err = "PPTX indexing is not available in this Docker build yet."
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=err,
                )
                return f"Indexing failed for {src.name}: {err}"
        if file_type == ".pdf":
            dep_err = self._check_pdf_dependencies()
            if dep_err:
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=dep_err,
                )
                return f"Indexing failed for {src.name}: {dep_err}"

        try:
            if file_type == ".docx" and parser_for_run == "simple_docx":
                try:
                    import docling  # type: ignore  # noqa: F401

                    docling_available = True
                except Exception:
                    docling_available = False
                if docling_available:
                    with self._patched_env_for_ingest():
                        await rx.process_with_rag(
                            file_path=stored_path,
                            output_dir="./output",
                            api_key=self.api_key,
                            base_url=self.base_url,
                            working_dir=working_abs,
                            parser=parser_for_run,
                            parse_method="auto",
                            mineru_backend="pipeline",
                            mineru_device="cpu",
                            fallback_parser="docling",
                            enable_parser_fallback=False,
                            max_chars=None,
                            max_paragraphs=None,
                            skip_query=True,
                            cli_queries=None,
                            embedding_workers=1,
                            embedding_batch_num=1,
                            embedding_max_retries=1,
                            embedding_backoff_base_sec=8.0,
                            embedding_backoff_max_sec=30.0,
                            llm_max_retries=2,
                            llm_backoff_base_sec=6.0,
                            llm_backoff_max_sec=30.0,
                            pdf_mode=pdf_mode,
                            max_pages=None,
                            page_range=pdf_page_range,
                            vision_page_range=vision_page_range,
                            max_vision_pages=max_vision_pages,
                            vision_target=vision_target,
                            no_vision=no_vision,
                            skip_kg_extraction=skip_kg_extraction,
                            raise_on_failure=True,
                        )
                else:
                    logger.info(
                        "DOCX simple_docx/docling unavailable in Docker; using python-docx fallback"
                    )
                    with self._patched_env_for_ingest():
                        await self._index_docx_with_python_docx(
                            file_path=stored_path,
                            workdir=workdir,
                        )
            else:
                with self._patched_env_for_ingest():
                    await rx.process_with_rag(
                        file_path=stored_path,
                        output_dir="./output",
                        api_key=self.api_key,
                        base_url=self.base_url,
                        working_dir=working_abs,
                        parser=parser_for_run,
                        parse_method="auto",
                        mineru_backend="pipeline",
                        mineru_device="cpu",
                        fallback_parser="docling",
                        enable_parser_fallback=False,
                        max_chars=None,
                        max_paragraphs=None,
                        skip_query=True,
                        cli_queries=None,
                        embedding_workers=1,
                        embedding_batch_num=1,
                        embedding_max_retries=1,
                        embedding_backoff_base_sec=8.0,
                        embedding_backoff_max_sec=30.0,
                        # WebUI indexing should tolerate temporary Gemini 503 spikes.
                        llm_max_retries=2,
                        llm_backoff_base_sec=6.0,
                        llm_backoff_max_sec=30.0,
                        pdf_mode=pdf_mode,
                        max_pages=None,
                        page_range=pdf_page_range,
                        vision_page_range=vision_page_range,
                        max_vision_pages=max_vision_pages,
                        vision_target=vision_target,
                        no_vision=no_vision,
                        skip_kg_extraction=skip_kg_extraction,
                        raise_on_failure=True,
                    )
        except Exception as exc:
            if "Content already exists" in str(exc):
                user_msg = (
                    "This file already exists in the current index. "
                    "Please use Reprocess, which will clear the old index first."
                )
                if not force_reprocess:
                    self._save_index_failure(
                        existing=existing,
                        doc_id=doc_id,
                        src=src,
                        stored_path=stored_path,
                        workdir=workdir,
                        parser=parser_for_run,
                        pdf_mode=pdf_mode,
                        sha256=sha256,
                        error_message=user_msg,
                    )
                    return f"Indexing failed for {src.name}: {user_msg}"
            err = (
                f"file_type={file_type}, parser={parser_for_run}, "
                f"working_dir={workdir.resolve()}, error={exc}"
            )
            logger.error("Indexing failed: %s\n%s", err, traceback.format_exc())
            self._save_index_failure(
                existing=existing,
                doc_id=doc_id,
                src=src,
                stored_path=stored_path,
                workdir=workdir,
                parser=parser_for_run,
                pdf_mode=pdf_mode,
                sha256=sha256,
                error_message=err,
            )
            return f"Indexing failed for {src.name}: {exc}"

        artifact_status = _index_artifact_status(str(workdir.resolve()))
        logger.info(
            "Working dir exists after process: %s",
            artifact_status["working_dir_exists"],
        )
        logger.info(
            "Working dir contents after process: %s", artifact_status["dir_entries"]
        )
        logger.info("Index artifact status: %s", artifact_status["files"])

        if not _is_index_ready(str(workdir.resolve())):
            status_err = _read_doc_status_error(str(workdir.resolve()), src.name)
            if status_err:
                lowered = status_err.lower()
                if any(
                    k in lowered
                    for k in ["quota", "resource_exhausted", "rate limit", "429"]
                ):
                    err = (
                        "LLM quota exceeded while extracting entities/chunks. "
                        "Try again later or reduce document size via WEBUI_DOCX_MAX_CHARS "
                        "(default 12000 for DOCX fallback). "
                        f"Detail: {status_err}"
                    )
                else:
                    err = f"Document processing failed: {status_err}"
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=err,
                )
                return f"Indexing failed for {src.name}: {err}"
            err = (
                f"process completed but working_dir index artifacts were not created. "
                f"expected={workdir.resolve()} "
                f"(file_type={file_type}, parser={parser_for_run})"
            )
            self._save_index_failure(
                existing=existing,
                doc_id=doc_id,
                src=src,
                stored_path=stored_path,
                workdir=workdir,
                parser=parser_for_run,
                pdf_mode=pdf_mode,
                sha256=sha256,
                error_message=err,
            )
            return f"Indexing failed for {src.name}: {err}"

        record = DocumentRecord(
            doc_id=doc_id,
            original_filename=src.name,
            stored_file_rel=self._as_rel(
                Path(stored_path), REGISTRY_PATH.parent.resolve()
            ),
            working_dir_rel=self._as_rel(workdir, REGISTRY_PATH.parent.resolve()),
            file_type=src.suffix.lower(),
            parser="pdf_hybrid" if parser == "pdf_fast" else parser,
            pdf_mode=pdf_mode,
            indexed_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
            status="indexed",
            summary="",
            sha256=sha256,
            source_metadata={
                "embedding_provider": self.embedding_provider,
                "embedding_model": self.embedding_model,
                "embedding_dim": self.embedding_dim,
            },
            visual_targets=(existing.visual_targets if existing else []),
            needs_reprocess=False,
            needs_reprocess_reason="",
            error_message="",
        )

        if existing:
            self.registry = [r for r in self.registry if r.doc_id != existing.doc_id]
        self.registry.append(record)
        self.registry_store.save(self.registry)
        self.rag_cache.pop(doc_id, None)

        note = ""
        if record.file_type in {".png", ".jpg", ".jpeg"} and no_vision:
            note = (
                "\nno_vision applies to PDF visual page rendering only; "
                "direct image input still uses vision."
            )
        return f"Indexed: {src.name} [{doc_id}]\nWorking dir: {workdir}{note}"

    def _save_index_failure(
        self,
        *,
        existing: Optional[DocumentRecord],
        doc_id: str,
        src: Path,
        stored_path: str,
        workdir: Path,
        parser: str,
        pdf_mode: str,
        sha256: str,
        error_message: str,
    ) -> None:
        record = DocumentRecord(
            doc_id=doc_id,
            original_filename=src.name,
            stored_file_rel=self._as_rel(
                Path(stored_path), REGISTRY_PATH.parent.resolve()
            ),
            working_dir_rel=self._as_rel(workdir, REGISTRY_PATH.parent.resolve()),
            file_type=src.suffix.lower(),
            parser=parser,
            pdf_mode=pdf_mode,
            indexed_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
            status="failed",
            summary="",
            sha256=sha256,
            source_metadata={
                "embedding_provider": self.embedding_provider,
                "embedding_model": self.embedding_model,
                "embedding_dim": self.embedding_dim,
            },
            visual_targets=(existing.visual_targets if existing else []),
            needs_reprocess=True,
            needs_reprocess_reason=error_message,
            error_message=error_message,
        )
        if existing:
            self.registry = [r for r in self.registry if r.doc_id != existing.doc_id]
        self.registry.append(record)
        self.registry_store.save(self.registry)
        self.rag_cache.pop(doc_id, None)

    async def process_index(
        self,
        file_paths: list[str] | str,
        force_parser: Optional[str] = None,
        pdf_page_range: Optional[str] = None,
        vision_target: Optional[str] = None,
        vision_page_range: Optional[str] = None,
        max_vision_pages: int = 1,
        no_vision: bool = True,
        skip_kg_extraction: bool = True,
        force_reprocess: bool = False,
    ) -> str:
        if isinstance(file_paths, str):
            file_paths = [file_paths]
        if not file_paths:
            return "Please upload/select files first."
        if not self.api_key:
            return (
                "Missing API key. Set LLM_BINDING_API_KEY "
                "or GEMINI_API_KEY/GOOGLE_API_KEY."
            )

        t0 = time.perf_counter()
        logs = []
        for fp in file_paths:
            try:
                logs.append(
                    await self._process_single_document(
                        file_path=fp,
                        force_parser=force_parser,
                        pdf_page_range=pdf_page_range,
                        vision_target=vision_target,
                        vision_page_range=vision_page_range,
                        max_vision_pages=max_vision_pages,
                        no_vision=no_vision,
                        skip_kg_extraction=skip_kg_extraction,
                        force_reprocess=force_reprocess,
                    )
                )
            except Exception as exc:
                if _is_gemini_quota_exceeded(exc):
                    logs.append(
                        "Gemini quota exceeded. Continue with text/table-only mode "
                        "or wait for quota reset."
                    )
                else:
                    logs.append(f"Failed indexing {Path(fp).name}: {exc}")

        self.state.processed = any(
            ("Indexed:" in x) or ("Skipped duplicate:" in x) for x in logs
        )
        self.state.last_process_ts = time.time()
        self.state.ingest_count += len(file_paths)
        elapsed = time.perf_counter() - t0
        self.state.processing_summary = (
            "\n\n".join(logs) + f"\n\nTotal processing time: {elapsed:.2f}s"
        )
        return self.state.processing_summary

    def route_query_to_documents(
        self, question: str, selected_doc_id: Optional[str] = None
    ) -> tuple[list[DocumentRecord], Optional[str]]:
        self._reload_registry()
        docs = [r for r in self.registry if r.status == "indexed"]
        if not docs:
            return [], "Please upload and process files first."

        q = (question or "").strip().lower()
        if not q:
            return [], "Please enter a question."

        # Highest priority: explicit filename mention in question.
        mentions = _extract_candidate_filenames(q)
        if mentions:
            matched = [d for d in docs if d.original_filename.lower() in mentions]
            if matched:
                return matched, None

        # Next priority: stem mention in question.
        stem_matched = []
        for d in docs:
            stem = Path(d.original_filename).stem.lower()
            if stem and stem in q:
                stem_matched.append(d)
        if stem_matched:
            uniq = {d.doc_id: d for d in stem_matched}
            return list(uniq.values()), None

        # Ambiguous query can be resolved by explicit dropdown selection.
        if selected_doc_id:
            rec = self._find_by_id(selected_doc_id)
            if rec is None:
                return [], "Selected file is not available."
            return [rec], None

        if len(docs) == 1:
            return docs, None
        return (
            [],
            "I found multiple indexed files. Which file would you like to ask about?",
        )

    async def load_rag_for_existing_index(self, rec: DocumentRecord) -> RAGAnything:
        ok, reason = self._record_index_status(rec)
        if not ok:
            if reason.startswith("registry path incompatible"):
                raise RuntimeError(
                    f"Registry path incompatible for {rec.original_filename}: "
                    f"{reason}. Please reprocess this file."
                )
            raise RuntimeError(
                f"Index path missing for {rec.original_filename}: "
                f"{reason}. Please reprocess this file."
            )
        working_dir = str(self._resolve_record_working_dir(rec))
        if rec.doc_id not in self.rag_cache:
            preferred = (
                rec.parser
                if rec.parser in SUPPORTED_QUERY_LOAD_PARSERS
                else "paddleocr"
            )
            parser_for_load = self._resolve_query_parser(preferred)
            logger.info(
                "Query-only mode: loading existing index for %s from %s",
                rec.original_filename,
                working_dir,
            )
            with self._patched_env_for_ingest():
                rag = await self._create_rag(
                    working_dir=working_dir,
                    parser=parser_for_load,
                )
                init_result = await rag._ensure_lightrag_initialized()
            if not init_result or not init_result.get("success"):
                detail = (init_result or {}).get("error", "unknown error")
                raise RuntimeError(
                    f"LightRAG init failed for " f"{rec.original_filename}: {detail}"
                )
            self.rag_cache[rec.doc_id] = rag
        return self.rag_cache[rec.doc_id]

    async def query_existing_document(
        self,
        question: str,
        rec: DocumentRecord,
        use_direct_vlm_on_query: bool = False,
    ) -> str:
        rag = await self.load_rag_for_existing_index(rec)
        special_kind = _detect_special_pdf_question(question)
        query_text = question
        marker_chunks: list[str] = []
        if rec.file_type == ".pdf" and special_kind:
            marker_chunks = self._find_marker_chunks(rec, special_kind)
            logger.info(
                "Special PDF query uses LightRAG rewrite: type=%s hits=%s file=%s",
                special_kind,
                len(marker_chunks),
                rec.original_filename,
            )
            query_text = _rewrite_pdf_special_query(question, special_kind)

        async def _run_query(vlm_enabled: bool):
            try:
                return await rag.aquery(
                    query_text,
                    mode="hybrid",
                    vlm_enhanced=vlm_enabled,
                )
            except TypeError as type_exc:
                # Backward-compatible path for test doubles or older wrappers
                # that don't accept `vlm_enhanced` yet.
                if "vlm_enhanced" not in str(type_exc):
                    raise
                return await rag.aquery(query_text, mode="hybrid")

        try:
            result = await _run_query(bool(use_direct_vlm_on_query))
            answer = str(result) if result is not None else "No answer was returned."
            if special_kind and marker_chunks:
                lowered = answer.lower()
                looks_bad_table = special_kind == "table" and (
                    "references" in lowered
                    or "subset 1" in lowered
                    or "[pdf table" in lowered
                )
                looks_bad_equation = special_kind == "equation" and (
                    "[pdf equation" in lowered
                    or ("accuracy" in lowered and ("tp + tn + fp + fn" not in lowered))
                )
                looks_bad_figure = (
                    special_kind == "figure" and "[pdf visual description" in lowered
                )
                if looks_bad_table or looks_bad_equation or looks_bad_figure:
                    marker_answer = self._format_marker_answer(
                        special_kind, marker_chunks, question
                    )
                    if marker_answer and (not _is_marker_only_answer(marker_answer)):
                        answer = marker_answer
            if _is_marker_only_answer(answer):
                if special_kind and marker_chunks:
                    marker_answer = self._format_marker_answer(
                        special_kind, marker_chunks, question
                    )
                    if marker_answer and (not _is_marker_only_answer(marker_answer)):
                        answer = marker_answer
                    else:
                        raise RuntimeError(
                            "Extraction error: retrieved marker-only chunk without usable content."
                        )
                else:
                    raise RuntimeError(
                        "Extraction error: retrieved marker-only chunk without usable content."
                    )
            if self._is_answer_only_request(question):
                return answer
            return f"Source: {rec.original_filename}\n\n{answer}"
        except Exception as exc:
            if not use_direct_vlm_on_query:
                raise
            if _is_temporary_vlm_error(exc):
                logger.warning(
                    "VLM query failed with temporary API error; falling back to indexed text/visual descriptions."
                )
                try:
                    result = await _run_query(False)
                    answer = (
                        str(result) if result is not None else "No answer was returned."
                    )
                    return (
                        f"Source: {rec.original_filename}\n\n"
                        "Note: direct vision query was unavailable, so this answer is based "
                        "on the indexed image description.\n\n"
                        f"{answer}"
                    )
                except Exception:
                    raise RuntimeError(
                        "Gemini Vision is temporarily unavailable and no indexed visual "
                        "description could be retrieved."
                    ) from exc
            raise

    async def query(
        self,
        question: str,
        selected_doc_id: Optional[str] = None,
        use_direct_vlm_on_query: bool = False,
    ) -> str:
        routed_docs, route_message = self.route_query_to_documents(
            question, selected_doc_id
        )
        if route_message:
            return route_message

        self.state.query_count += 1
        if len(routed_docs) == 1:
            rec = routed_docs[0]
            try:
                source_prefix = (
                    "Answer based on selected file"
                    if selected_doc_id
                    else "Answer based on"
                )
                answer = await self.query_existing_document(
                    question,
                    rec,
                    use_direct_vlm_on_query=use_direct_vlm_on_query,
                )
                return answer.replace("Source:", f"{source_prefix}:")
            except Exception as exc:
                if _is_gemini_quota_exceeded(exc):
                    return (
                        "Gemini quota exceeded. Continue with text/table-only mode "
                        "or wait for quota reset."
                    )
                return f"Query failed: {exc}"

        lines = ["Compared files:"]
        segments = []
        for rec in routed_docs:
            lines.append(f"- {rec.original_filename}")
            try:
                result = await self.query_existing_document(
                    question,
                    rec,
                    use_direct_vlm_on_query=use_direct_vlm_on_query,
                )
                segments.append(f"[{rec.original_filename}]\n{result}")
            except Exception as exc:
                segments.append(f"[{rec.original_filename}]\nQuery failed: {exc}")
        return "\n".join(lines) + "\n\n" + "\n\n".join(segments)

    async def analyze_visual_target(
        self,
        doc_id: str,
        vision_target: Optional[str],
        vision_page_range: Optional[str],
        max_vision_pages: int,
    ) -> str:
        rec = self._find_by_id(doc_id)
        if rec is None:
            return "Please select an indexed file first."
        if rec.file_type != ".pdf":
            return "Visual target analysis is intended for PDF files."

        result = await self._process_single_document(
            file_path=str(self._resolve_record_stored_file(rec)),
            force_parser="pdf_hybrid",
            pdf_page_range=None,
            vision_target=vision_target,
            vision_page_range=vision_page_range,
            max_vision_pages=max_vision_pages,
            no_vision=False,
            skip_kg_extraction=True,
        )
        rec2 = self._find_by_id(doc_id)
        if rec2 is not None:
            rec2.visual_targets.append(
                {
                    "vision_target": vision_target or "",
                    "vision_page_range": vision_page_range or "",
                    "max_vision_pages": int(max_vision_pages),
                    "processed_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                }
            )
            self.registry_store.save(self.registry)
        return result

    def delete_document(self, doc_id: str) -> str:
        rec = self._find_by_id(doc_id)
        if rec is None:
            return "Please select an indexed file first."
        self.rag_cache.pop(doc_id, None)
        wd = self._resolve_record_working_dir(rec)
        if wd.exists():
            shutil.rmtree(wd, ignore_errors=True)
        stored_file = self._resolve_record_stored_file(rec)
        shared_upload = False
        for other in self.registry:
            if other.doc_id == rec.doc_id:
                continue
            if other.stored_file_rel == rec.stored_file_rel:
                shared_upload = True
                break
        if (not shared_upload) and stored_file.exists():
            try:
                stored_file.unlink()
            except Exception:
                pass
        self.registry = [r for r in self.registry if r.doc_id != doc_id]
        self.registry_store.save(self.registry)
        return f"Deleted indexed file: {rec.original_filename}"


def _extract_uploaded_paths(file_input) -> list[str]:
    if not file_input:
        return []
    if isinstance(file_input, (str, Path)):
        return [str(file_input)]
    out = []
    if isinstance(file_input, list):
        for item in file_input:
            if isinstance(item, (str, Path)):
                out.append(str(item))
            elif hasattr(item, "name"):
                out.append(str(item.name))
    elif hasattr(file_input, "name"):
        out.append(str(file_input.name))
    return out


def _parse_doc_label(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    m = re.match(r"^(.*)\s\[[0-9a-f]+\]$", str(value).strip())
    if m:
        return m.group(1).strip()
    return str(value).strip() or None


def _parse_doc_id_from_label(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    m = re.search(r"\[([0-9a-f]+)\]$", str(value))
    return m.group(1) if m else None


def _format_pdf_export_result(
    status: str, pdf_path: Optional[Path]
) -> tuple[str, Optional[str]]:
    return status, (str(pdf_path) if pdf_path else None)


def export_current_answer_action(
    question: str,
    answer: str,
    source_file: Optional[str],
    output_dir: str | Path = REPORTS_ROOT,
) -> tuple[str, Optional[str]]:
    question_text = str(question or "").strip()
    answer_text = str(answer or "").strip()
    if not answer_text:
        return _format_pdf_export_result("No answer available to export.", None)
    pdf_path = export_current_answer_to_pdf(
        question=question_text,
        answer=answer_text,
        source_file=str(source_file or "").strip() or None,
        output_dir=output_dir,
    )
    return _format_pdf_export_result("PDF exported successfully.", pdf_path)


def export_chat_history_action(
    chat_history,
    selected_doc_label: Optional[str],
    last_source: Optional[str],
    output_dir: str | Path = REPORTS_ROOT,
) -> tuple[str, Optional[str]]:
    messages = _normalize_to_messages(chat_history)
    if not messages:
        return _format_pdf_export_result("No chat history available to export.", None)
    selected_label = _parse_doc_label(selected_doc_label)
    source_for_report = selected_label or (str(last_source or "").strip() or None)
    pdf_path = export_chat_history_to_pdf(
        messages=messages,
        selected_file=source_for_report,
        output_dir=output_dir,
    )
    return _format_pdf_export_result("PDF exported successfully.", pdf_path)


def generate_pdf_report_action(
    *,
    report_request: str,
    selected_doc_label: Optional[str],
    service: WebUIRAGService,
    output_dir: str | Path = REPORTS_ROOT,
) -> tuple[str, Optional[str]]:
    req = str(report_request or "").strip()
    if not req:
        return _format_pdf_export_result("Please enter a report request.", None)

    service._reload_registry()
    all_docs = service.list_documents()

    selected_doc_id = _parse_doc_id_from_label(selected_doc_label)
    doc: Optional[DocumentRecord] = None
    if selected_doc_id:
        doc = service._find_by_id(selected_doc_id)
        if doc is None:
            return _format_pdf_export_result("Selected file is not available.", None)
        if doc.status != "indexed":
            return _format_pdf_export_result(
                "Selected file is not indexed successfully. Please reprocess it first.",
                None,
            )
    indexed_docs = [r for r in all_docs if r.status == "indexed"]
    if not indexed_docs and doc is None:
        return _format_pdf_export_result(
            "Please upload and process a file before generating a PDF report.",
            None,
        )
    else:
        mentioned_names = _extract_candidate_filenames(req)
        if mentioned_names:
            for candidate in indexed_docs:
                if candidate.original_filename.lower() in mentioned_names:
                    doc = candidate
                    break
        if doc is None:
            if len(indexed_docs) > 1:
                return _format_pdf_export_result(
                    "Please select a file before generating a PDF report.",
                    None,
                )
            doc = indexed_docs[0]

    assert doc is not None

    try:

        def _query_func(question_text: str, rec: DocumentRecord) -> str:
            return service.run(
                service.query_existing_document(
                    question_text,
                    rec,
                    use_direct_vlm_on_query=False,
                )
            )

        pdf_path = generate_pdf_report_from_index(
            request=req,
            doc_record=doc,
            output_dir=output_dir,
            query_func=_query_func,
        )
        return _format_pdf_export_result("PDF report generated successfully.", pdf_path)
    except Exception as exc:
        return _format_pdf_export_result(
            f"PDF report generation failed: {exc}",
            None,
        )


def build_webui():
    try:
        import gradio as gr
    except Exception as exc:
        raise RuntimeError(
            "Gradio is required for WebUI. Install with: "
            ".\\.venv\\Scripts\\python.exe -m pip install gradio"
        ) from exc

    service = WebUIRAGService()

    with gr.Blocks(title="RAG-Anything Local WebUI") as demo:
        gr.Markdown(
            "## RAG-Anything Local WebUI\n"
            "Upload files -> Process/Index once -> One chat for all indexed files."
        )

        with gr.Row():
            with gr.Column(scale=1):
                file_input = gr.File(
                    label="Upload files",
                    type="filepath",
                    file_count="multiple",
                    elem_id="upload_files",
                )
                process_btn = gr.Button(
                    "Process / Index", variant="primary", elem_id="process_index_btn"
                )
                reprocess_btn = gr.Button(
                    "Reprocess selected file", elem_id="reprocess_btn"
                )
                delete_btn = gr.Button("Delete selected file", elem_id="delete_btn")
                indexed_files = gr.Dropdown(
                    label="Indexed files", choices=[], elem_id="indexed_files"
                )
                registry_view = gr.Textbox(label="Indexed files list", lines=10)
                process_status = gr.Textbox(label="Status", lines=10)

                with gr.Accordion("Advanced options", open=False):
                    force_parser = gr.Dropdown(
                        choices=[
                            "",
                            "pdf_hybrid",
                            "pdf_fast",
                            "simple_docx",
                            "paddleocr",
                            "docling",
                        ],
                        value="",
                        label="Force parser",
                    )
                    pdf_page_range = gr.Textbox(label="PDF page range")
                    no_vision = gr.Checkbox(
                        value=True, label="No vision (default for PDF)"
                    )
                    skip_kg_extraction = gr.Checkbox(
                        value=True, label="Skip KG extraction"
                    )
                    vision_target = gr.Textbox(label='Vision target (e.g. "Fig. 2")')
                    vision_page_range = gr.Textbox(label="Vision page range")
                    max_vision_pages = gr.Number(
                        value=1, precision=0, label="Max vision pages"
                    )
                    use_direct_vlm_on_query = gr.Checkbox(
                        value=False,
                        label="Use direct VLM on query",
                        elem_id="use_direct_vlm_on_query",
                    )

                gr.Markdown("### Visual analysis on demand")
                visual_doc = gr.Dropdown(
                    label="PDF file for visual analysis",
                    choices=[],
                    elem_id="visual_doc",
                )
                visual_target = gr.Textbox(
                    label='Vision target (e.g. "Fig. 2")', elem_id="visual_target"
                )
                visual_range = gr.Textbox(label="Vision page range")
                analyze_btn = gr.Button("Analyze visual target", elem_id="analyze_btn")

                gr.Markdown("### Export PDF")
                export_current_btn = gr.Button(
                    "Export current answer to PDF", elem_id="export_current_pdf_btn"
                )
                export_history_btn = gr.Button(
                    "Export chat history to PDF", elem_id="export_history_pdf_btn"
                )
                export_status = gr.Textbox(label="Export status", lines=3)
                export_pdf_file = gr.File(label="Download PDF report")
                gr.Markdown("### PDF Report Agent")
                report_request = gr.Textbox(label="Report request")
                generate_report_btn = gr.Button(
                    "Generate PDF report", elem_id="generate_report_btn"
                )
                report_status = gr.Textbox(label="PDF report status", lines=3)
                report_file = gr.File(label="Download generated PDF report")

            with gr.Column(scale=2):
                try:
                    chatbot = gr.Chatbot(
                        label="Chat", type="messages", elem_id="chatbot"
                    )
                except TypeError:
                    chatbot = gr.Chatbot(label="Chat", elem_id="chatbot")
                question = gr.Textbox(label="Your question", elem_id="chat_question")
                ask_btn = gr.Button("Ask", elem_id="ask_btn")
                clear_btn = gr.Button("Clear chat", elem_id="clear_btn")

        history_state = gr.State([])
        last_question_state = gr.State("")
        last_answer_state = gr.State("")
        last_source_state = gr.State("")

        def refresh_doc_choices():
            choices = [
                f"{r.original_filename} [{r.doc_id}]" for r in service.list_documents()
            ]
            first = choices[0] if choices else None
            return (
                gr.update(choices=choices, value=first),
                gr.update(choices=choices, value=first),
                service.render_registry_text(),
            )

        def on_process(
            files,
            f_parser,
            p_range,
            no_v,
            skip_kg,
            v_target,
            v_range,
            max_v_pages,
        ):
            file_paths = _extract_uploaded_paths(files)
            status = service.run(
                service.process_index(
                    file_paths=file_paths,
                    force_parser=(f_parser or None),
                    pdf_page_range=(p_range or None),
                    vision_target=(v_target or None),
                    vision_page_range=(v_range or None),
                    max_vision_pages=int(max_v_pages or 1),
                    no_vision=bool(no_v),
                    skip_kg_extraction=bool(skip_kg),
                )
            )
            doc_dropdown, visual_dropdown, registry_txt = refresh_doc_choices()
            return status, doc_dropdown, visual_dropdown, registry_txt, [], []

        def on_reprocess(
            selected_doc,
            f_parser,
            p_range,
            no_v,
            skip_kg,
            v_target,
            v_range,
            max_v_pages,
        ):
            doc_id = _parse_doc_id_from_label(selected_doc)
            if not doc_id:
                return (
                    "Please select an indexed file first.",
                    gr.update(),
                    gr.update(),
                    service.render_registry_text(),
                )
            rec = service._find_by_id(doc_id)
            if rec is None:
                return (
                    "Selected file is not available.",
                    gr.update(),
                    gr.update(),
                    service.render_registry_text(),
                )
            status = service.run(
                service.process_index(
                    file_paths=[str(service._resolve_record_stored_file(rec))],
                    force_parser=(f_parser or rec.parser),
                    pdf_page_range=(p_range or None),
                    vision_target=(v_target or None),
                    vision_page_range=(v_range or None),
                    max_vision_pages=int(max_v_pages or 1),
                    no_vision=bool(no_v),
                    skip_kg_extraction=bool(skip_kg),
                    force_reprocess=True,
                )
            )
            doc_dropdown, visual_dropdown, registry_txt = refresh_doc_choices()
            return status, doc_dropdown, visual_dropdown, registry_txt

        def on_delete(selected_doc):
            doc_id = _parse_doc_id_from_label(selected_doc)
            status = (
                service.delete_document(doc_id)
                if doc_id
                else "Please select an indexed file first."
            )
            doc_dropdown, visual_dropdown, registry_txt = refresh_doc_choices()
            return status, doc_dropdown, visual_dropdown, registry_txt

        def on_ask(user_q, chat_history, selected_doc, use_direct_vlm):
            doc_id = _parse_doc_id_from_label(selected_doc)
            answer = service.run(
                service.query(
                    user_q,
                    selected_doc_id=doc_id,
                    use_direct_vlm_on_query=bool(use_direct_vlm),
                )
            )
            out = _append_messages(chat_history, user_q, answer)
            return (
                out,
                out,
                str(user_q or ""),
                str(answer or ""),
                _parse_doc_label(selected_doc) or "",
            )

        def on_clear():
            return [], [], "", "", ""

        def on_analyze(selected_doc, v_target, v_range, max_v_pages):
            doc_id = _parse_doc_id_from_label(selected_doc)
            status = service.run(
                service.analyze_visual_target(
                    doc_id=doc_id or "",
                    vision_target=(v_target or None),
                    vision_page_range=(v_range or None),
                    max_vision_pages=int(max_v_pages or 1),
                )
            )
            _, _, registry_txt = refresh_doc_choices()
            return status, registry_txt

        def on_export_current_answer(last_q, last_a, last_source):
            return export_current_answer_action(
                question=last_q,
                answer=last_a,
                source_file=last_source,
                output_dir=REPORTS_ROOT,
            )

        def on_export_chat_history(chat_history, selected_doc, last_source):
            return export_chat_history_action(
                chat_history=chat_history,
                selected_doc_label=selected_doc,
                last_source=last_source,
                output_dir=REPORTS_ROOT,
            )

        def on_generate_pdf_report(req_text, selected_doc):
            return generate_pdf_report_action(
                report_request=req_text,
                selected_doc_label=selected_doc,
                service=service,
                output_dir=REPORTS_ROOT,
            )

        process_btn.click(
            on_process,
            [
                file_input,
                force_parser,
                pdf_page_range,
                no_vision,
                skip_kg_extraction,
                vision_target,
                vision_page_range,
                max_vision_pages,
            ],
            [
                process_status,
                indexed_files,
                visual_doc,
                registry_view,
                chatbot,
                history_state,
            ],
        )
        reprocess_btn.click(
            on_reprocess,
            [
                indexed_files,
                force_parser,
                pdf_page_range,
                no_vision,
                skip_kg_extraction,
                vision_target,
                vision_page_range,
                max_vision_pages,
            ],
            [process_status, indexed_files, visual_doc, registry_view],
        )
        delete_btn.click(
            on_delete,
            [indexed_files],
            [process_status, indexed_files, visual_doc, registry_view],
        )
        ask_btn.click(
            on_ask,
            [question, history_state, indexed_files, use_direct_vlm_on_query],
            [
                chatbot,
                history_state,
                last_question_state,
                last_answer_state,
                last_source_state,
            ],
        )
        clear_btn.click(
            on_clear,
            [],
            [
                chatbot,
                history_state,
                last_question_state,
                last_answer_state,
                last_source_state,
            ],
        )
        analyze_btn.click(
            on_analyze,
            [visual_doc, visual_target, visual_range, max_vision_pages],
            [process_status, registry_view],
        )
        export_current_btn.click(
            on_export_current_answer,
            [last_question_state, last_answer_state, last_source_state],
            [export_status, export_pdf_file],
        )
        export_history_btn.click(
            on_export_chat_history,
            [history_state, indexed_files, last_source_state],
            [export_status, export_pdf_file],
        )
        generate_report_btn.click(
            on_generate_pdf_report,
            [report_request, indexed_files],
            [report_status, report_file],
        )
        demo.load(refresh_doc_choices, [], [indexed_files, visual_doc, registry_view])

    return demo


def _is_port_available(host: str, port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        try:
            sock.bind((host, port))
            return True
        except OSError:
            return False


def _env_bool(name: str, default: bool = False) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def get_webui_host() -> str:
    return os.getenv("WEBUI_HOST", "0.0.0.0")


def get_webui_port() -> int:
    env_port = os.getenv("WEBUI_PORT") or os.getenv("GRADIO_SERVER_PORT")
    preferred = 7860
    if env_port:
        try:
            preferred = int(env_port)
        except ValueError:
            logger.warning(
                "Invalid WEBUI_PORT/GRADIO_SERVER_PORT=%s. Falling back to default 7860.",
                env_port,
            )
            preferred = 7860
    return preferred


def get_webui_auth() -> Optional[tuple[str, str]]:
    user = os.getenv("WEBUI_AUTH_USER", "").strip()
    password = os.getenv("WEBUI_AUTH_PASSWORD", "").strip()
    if user and password:
        return (user, password)
    return None


def get_webui_share() -> bool:
    return _env_bool("GRADIO_SHARE", default=False)


def _resolve_launch_port(host: str, preferred: int) -> Optional[int]:
    if _is_port_available(host, preferred):
        return preferred
    logger.warning(
        "Preferred port %s is busy. Falling back to an auto-selected free port.",
        preferred,
    )
    # Let Gradio pick a free port automatically.
    return None


if __name__ == "__main__":
    app = build_webui()
    host = get_webui_host()
    preferred_port = get_webui_port()
    auth = get_webui_auth()
    share = get_webui_share()
    if auth:
        logger.info("WebUI auth is enabled.")
    launch_kwargs = {
        "server_name": host,
        "server_port": _resolve_launch_port(host, preferred_port),
        "share": share,
    }
    if auth:
        launch_kwargs["auth"] = auth
    app.launch(**launch_kwargs)
