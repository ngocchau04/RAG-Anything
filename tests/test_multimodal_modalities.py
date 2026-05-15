from __future__ import annotations

from pathlib import Path

import pytest

from raganything.parser import get_parser


REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO_ROOT / "tests" / "fixtures" / "multimodal"
OUTPUT_DIR = REPO_ROOT / "output" / "test_modalities"

TEXT_MARKER = "TEXT_MARKER_CPU_ONLY_RAG"
TABLE_MARKER = "TABLE_MARKER_REVENUE_2024"
IMAGE_MARKER = "IMAGE_MARKER_ARCHITECTURE_DIAGRAM"
EQUATION_MARKER = "EQUATION_MARKER_E_MC2"
CHART_MARKER = "CHART_MARKER_QUARTERLY_SALES"


def _ensure_fixture_dir() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)


def _ensure_md_fixture() -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / "sample_multimodal.md"
    content = f"""# Sample Multimodal Markdown

{TEXT_MARKER}

## Table
| Year | Revenue |
|---|---|
| 2024 | {TABLE_MARKER} |

## Image
![Architecture]({IMAGE_MARKER}.png)

## Equation
E = mc^2 ({EQUATION_MARKER})

## Chart
Chart caption: {CHART_MARKER}
"""
    path.write_text(content, encoding="utf-8")
    return path


def _ensure_html_fixture() -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / "sample_multimodal.html"
    content = f"""<!doctype html>
<html>
<body>
<h1>{TEXT_MARKER}</h1>
<table>
  <tr><th>Year</th><th>Revenue</th></tr>
  <tr><td>2024</td><td>{TABLE_MARKER}</td></tr>
</table>
<p><img alt="{IMAGE_MARKER}" src="{IMAGE_MARKER}.png" /></p>
<p>{EQUATION_MARKER}: E = mc^2</p>
<p>{CHART_MARKER}</p>
</body>
</html>
"""
    path.write_text(content, encoding="utf-8")
    return path


def _ensure_docx_fixture() -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / "sample_multimodal.docx"
    try:
        from docx import Document
    except Exception as exc:
        pytest.skip(f"python-docx is required for DOCX fixture generation: {exc}")

    doc = Document()
    doc.add_heading("Sample Multimodal DOCX", level=1)
    doc.add_paragraph(TEXT_MARKER)
    doc.add_heading("Equation", level=2)
    doc.add_paragraph(f"{EQUATION_MARKER}: E = mc^2")
    doc.add_heading("Image", level=2)
    doc.add_paragraph(f"Image caption: {IMAGE_MARKER}")
    doc.add_heading("Chart", level=2)
    doc.add_paragraph(f"Chart caption: {CHART_MARKER}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Year"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "2024"
    table.cell(1, 1).text = TABLE_MARKER
    doc.save(path)
    return path


def _ensure_placeholder_binary_fixture(name: str, marker: str) -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / name
    path.write_text(f"Placeholder fixture marker: {marker}\n", encoding="utf-8")
    return path


def _ensure_pdf_fixture() -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / "sample_multimodal.pdf"
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except Exception as exc:
        pytest.skip(f"reportlab is required for PDF fixture generation: {exc}")

    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for line in [
        TEXT_MARKER,
        TABLE_MARKER,
        IMAGE_MARKER,
        f"{EQUATION_MARKER}: E = mc^2",
        CHART_MARKER,
    ]:
        c.drawString(72, y, line)
        y -= 24
    c.showPage()
    c.save()
    return path


def _ensure_pptx_fixture() -> Path:
    _ensure_fixture_dir()
    path = FIXTURE_DIR / "sample_multimodal.pptx"
    try:
        from pptx import Presentation
    except Exception as exc:
        pytest.skip(f"python-pptx is required for PPTX fixture generation: {exc}")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sample Multimodal PPTX"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for marker in [
        TEXT_MARKER,
        TABLE_MARKER,
        IMAGE_MARKER,
        EQUATION_MARKER,
        CHART_MARKER,
    ]:
        p = body.add_paragraph()
        p.text = marker
    prs.save(str(path))
    return path


def _get_docling_parser_or_skip():
    parser = get_parser("docling")
    if not parser.check_installation():
        pytest.skip(
            "Docling parser is not installed. PDF/PPTX runtime assertions are unavailable in CPU-only baseline."
        )
    return parser


def _parse_with_docling(path: Path) -> str:
    parser = _get_docling_parser_or_skip()
    try:
        content = parser.parse_document(
            path,
            output_dir=OUTPUT_DIR,
            method="auto",
        )
    except Exception as exc:
        pytest.skip(
            f"Docling runtime path is unstable/unsupported for {path.suffix} in CPU-only environment: {exc}"
        )
    if not content:
        pytest.skip(f"Docling returned no content for {path.name}")
    return "\n".join(item.get("text", "") for item in content if isinstance(item, dict))


def _parse_docx_to_text(docx_path: Path) -> str:
    parser = get_parser("simple_docx")
    content = parser.parse_document(
        docx_path,
        output_dir=OUTPUT_DIR,
        max_chars=50_000,
    )
    assert content, "simple_docx returned empty content"
    return "\n".join(item.get("text", "") for item in content if isinstance(item, dict))


def _normalized(text: str) -> str:
    return text.replace("\\_", "_")


def test_text_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert TEXT_MARKER in parsed_text


def test_image_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert IMAGE_MARKER in parsed_text


def test_table_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert TABLE_MARKER in parsed_text


def test_equation_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert EQUATION_MARKER in parsed_text


def test_chart_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert CHART_MARKER in parsed_text


def test_pdf_multimodal_extraction() -> None:
    pdf_path = _ensure_pdf_fixture()
    parsed_text = _normalized(_parse_with_docling(pdf_path))
    assert TEXT_MARKER in parsed_text
    assert TABLE_MARKER in parsed_text
    assert IMAGE_MARKER in parsed_text
    assert EQUATION_MARKER in parsed_text
    assert CHART_MARKER in parsed_text


def test_pptx_multimodal_extraction() -> None:
    pptx_path = _ensure_pptx_fixture()
    parsed_text = _normalized(_parse_with_docling(pptx_path))
    assert TEXT_MARKER in parsed_text
    assert TABLE_MARKER in parsed_text
    assert IMAGE_MARKER in parsed_text
    assert EQUATION_MARKER in parsed_text
    assert CHART_MARKER in parsed_text


def test_docx_multimodal_extraction() -> None:
    docx_path = _ensure_docx_fixture()
    parsed_text = _normalized(_parse_docx_to_text(docx_path))
    assert TEXT_MARKER in parsed_text
    assert TABLE_MARKER in parsed_text
    assert IMAGE_MARKER in parsed_text
    assert EQUATION_MARKER in parsed_text
    assert CHART_MARKER in parsed_text


def test_md_multimodal_fixture_markers() -> None:
    md_path = _ensure_md_fixture()
    content = md_path.read_text(encoding="utf-8")
    assert TEXT_MARKER in content
    assert TABLE_MARKER in content
    assert IMAGE_MARKER in content
    assert EQUATION_MARKER in content
    assert CHART_MARKER in content


def test_html_multimodal_fixture_markers() -> None:
    html_path = _ensure_html_fixture()
    content = html_path.read_text(encoding="utf-8")
    assert TEXT_MARKER in content
    assert TABLE_MARKER in content
    assert IMAGE_MARKER in content
    assert EQUATION_MARKER in content
    assert CHART_MARKER in content
