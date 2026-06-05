from __future__ import annotations

import hashlib
import json
import os
import shutil
import time
import traceback
import uuid
from contextlib import contextmanager
from functools import partial
from pathlib import Path
from typing import Any, Optional

import numpy as np
from lightrag.utils import EmbeddingFunc, logger
from openai import AsyncOpenAI

from raganything import RAGAnything, RAGAnythingConfig
from raganything.config import resolve_embedding_runtime_config

from backend.app.schemas.document import UIState
from backend.app.schemas.document import DocumentRecord
from backend.app.services.document_service import (
    DocumentRegistryService,
    is_index_ready,
)


def safe_filename(name: str) -> str:
    import re

    out = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return out or "uploaded_file"


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def detect_parser_for_file(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return "pdf_hybrid"
    if suffix == ".docx":
        return "simple_docx"
    if suffix in {".png", ".jpg", ".jpeg"}:
        return "paddleocr"
    if suffix == ".pptx":
        return "docling"
    if suffix in {".md", ".html", ".htm"}:
        return "docling"
    return "docling"


def index_artifact_status(working_dir: str) -> dict[str, Any]:
    wd = Path(working_dir)
    files = {
        "graph_chunk_entity_relation.graphml": wd
        / "graph_chunk_entity_relation.graphml",
        "vdb_chunks.json": wd / "vdb_chunks.json",
        "kv_store_text_chunks.json": wd / "kv_store_text_chunks.json",
    }
    status = {
        "working_dir_exists": wd.exists(),
        "working_dir_abs": str(wd.resolve()),
        "files": {},
        "dir_entries": [],
    }
    for name, path in files.items():
        exists = path.exists()
        size = path.stat().st_size if exists else 0
        status["files"][name] = {"exists": exists, "size": size}
    if wd.exists():
        try:
            status["dir_entries"] = sorted([p.name for p in wd.iterdir()])[:80]
        except Exception:
            status["dir_entries"] = []
    return status


def read_doc_status_error(working_dir: str, filename: str) -> Optional[str]:
    status_path = Path(working_dir) / "kv_store_doc_status.json"
    if not status_path.exists():
        return None
    try:
        data = json.loads(status_path.read_text(encoding="utf-8"))
    except Exception:
        return None
    target = (filename or "").lower()
    for item in (data or {}).values():
        if not isinstance(item, dict):
            continue
        file_ref = str(item.get("file_path", "")).lower()
        if target and target not in file_ref:
            continue
        status = str(item.get("status", "")).lower()
        if status in {"failed", "error"}:
            return (
                str(item.get("error_msg", "") or "").strip() or "document status failed"
            )
    return None


async def ollama_embed_runtime(texts, *, model: str, host: str, target_dim: int):
    import aiohttp

    url = f"{host.rstrip('/')}/api/embed"
    async with aiohttp.ClientSession() as session:
        async with session.post(
            url,
            json={"model": model, "input": texts},
            timeout=60,
        ) as resp:
            if resp.status != 200:
                body = await resp.text()
                raise RuntimeError(f"Ollama embed failed ({resp.status}): {body[:240]}")
            data = await resp.json()
    embeddings = data.get("embeddings", [])
    if not embeddings:
        raise RuntimeError("Ollama returned empty embeddings.")
    if len(embeddings[0]) != target_dim:
        raise RuntimeError(
            f"Embedding dim mismatch. Expected {target_dim}, got {len(embeddings[0])}."
        )
    return np.array(embeddings, dtype=np.float32)


async def llm_call_runtime(
    prompt,
    *,
    model: str,
    api_key: Optional[str],
    base_url: Optional[str],
    system_prompt=None,
    history_messages=None,
    **kwargs,
):
    history_messages = history_messages or []
    messages = kwargs.pop("messages", None)
    kwargs.pop("hashing_kv", None)
    kwargs.pop("keyword_extraction", None)
    kwargs.pop("enable_cot", None)
    kwargs.pop("response_format", None)

    if messages is None:
        messages = []
        if system_prompt:
            messages.append({"role": "system", "content": system_prompt})
        messages.extend(history_messages)
        messages.append({"role": "user", "content": prompt})

    client = AsyncOpenAI(
        api_key=api_key,
        base_url=base_url,
        max_retries=0,
        timeout=180.0,
    )
    async with client:
        resp = await client.chat.completions.create(
            model=model,
            messages=messages,
            **kwargs,
        )
    content = resp.choices[0].message.content if resp.choices else None
    if content is None:
        raise RuntimeError("Empty LLM response content.")
    return content


async def vision_call_runtime(
    prompt,
    *,
    vision_model: str,
    api_key: Optional[str],
    base_url: Optional[str],
    llm_model: str,
    system_prompt=None,
    history_messages=None,
    image_data=None,
    messages=None,
    **kwargs,
):
    history_messages = history_messages or []
    if messages:
        return await llm_call_runtime(
            "",
            model=vision_model,
            api_key=api_key,
            base_url=base_url,
            messages=messages,
            **kwargs,
        )

    if image_data:
        payload = [
            {"role": "system", "content": system_prompt} if system_prompt else None,
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/jpeg;base64,{image_data}"},
                    },
                ],
            },
        ]
        payload = [m for m in payload if m]
        return await llm_call_runtime(
            "",
            model=vision_model,
            api_key=api_key,
            base_url=base_url,
            messages=payload,
            **kwargs,
        )

    return await llm_call_runtime(
        prompt,
        model=llm_model,
        api_key=api_key,
        base_url=base_url,
        system_prompt=system_prompt,
        history_messages=history_messages,
        **kwargs,
    )


class IndexingServiceMixin:
    @staticmethod
    def _check_pdf_dependencies() -> Optional[str]:
        missing = []
        try:
            import fitz  # type: ignore  # noqa: F401
        except Exception:
            missing.append("pymupdf(fitz)")
        try:
            import pdfplumber  # type: ignore  # noqa: F401
        except Exception:
            missing.append("pdfplumber")
        if missing:
            return (
                "PDF hybrid requires PyMuPDF and/or pdfplumber in Docker. "
                f"Missing: {', '.join(missing)}"
            )
        return None

    async def _index_docx_with_python_docx(
        self,
        file_path: str,
        workdir: Path,
    ) -> None:
        try:
            from docx import Document  # type: ignore
        except Exception as exc:
            raise RuntimeError(
                "simple_docx requires docling, and python-docx fallback is unavailable. "
                "Install docling or python-docx in Docker."
            ) from exc

        doc = Document(file_path)
        lines: list[str] = []
        for p in doc.paragraphs:
            txt = (p.text or "").strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                row_text = " | ".join([c for c in cells if c])
                if row_text:
                    lines.append(row_text)
        content = "\n".join(lines).strip()
        if not content:
            raise RuntimeError("DOCX extracted no text content.")
        max_chars_raw = os.getenv("WEBUI_DOCX_MAX_CHARS", "12000").strip()
        try:
            max_chars = int(max_chars_raw)
        except Exception:
            max_chars = 12000
        if max_chars > 0 and len(content) > max_chars:
            logger.info(
                "DOCX fallback truncation enabled: %s -> %s chars (set WEBUI_DOCX_MAX_CHARS to adjust)",
                len(content),
                max_chars,
            )
            content = content[:max_chars]

        rag = await self._create_rag(
            working_dir=str(workdir.resolve()), parser="paddleocr"
        )
        init_result = await rag._ensure_lightrag_initialized()
        if not init_result or not init_result.get("success"):
            detail = (init_result or {}).get("error", "unknown error")
            raise RuntimeError(f"LightRAG init failed for DOCX fallback: {detail}")
        await rag.insert_content_list(
            [{"type": "text", "text": content, "page_idx": 0}], file_path=file_path
        )

    def _persist_upload_file(self, file_path: str, doc_id: str) -> str:
        self.paths.uploads_dir.mkdir(parents=True, exist_ok=True)
        src = Path(file_path)
        dst = self.paths.uploads_dir / f"{doc_id}_{safe_filename(src.name)}"
        try:
            if src.resolve() == dst.resolve():
                return str(dst.resolve())
        except Exception:
            pass
        shutil.copy2(src, dst)
        return str(dst.resolve())

    @contextmanager
    def _patched_env_for_ingest(self):
        keys = {
            "EMBEDDING_PROVIDER": "ollama",
            "EMBEDDING_BINDING": "ollama",
            "EMBEDDING_MODEL": self.embedding_model,
            "EMBEDDING_DIM": str(self.embedding_dim),
            "OLLAMA_HOST": self.ollama_host,
        }
        old = {k: os.environ.get(k) for k in keys}
        try:
            for k, v in keys.items():
                os.environ[k] = v
            yield
        finally:
            for k, old_v in old.items():
                if old_v is None:
                    os.environ.pop(k, None)
                else:
                    os.environ[k] = old_v

    async def _ensure_ollama_model_available(self) -> None:
        import aiohttp

        if self.embedding_provider != "ollama":
            return
        url = f"{self.ollama_host.rstrip('/')}/api/tags"
        async with aiohttp.ClientSession() as session:
            async with session.get(url, timeout=20) as resp:
                if resp.status != 200:
                    raise RuntimeError(
                        f"Ollama host check failed ({resp.status}) at {url}"
                    )
                data = await resp.json()
        names = [
            m.get("name", "") for m in data.get("models", []) if isinstance(m, dict)
        ]
        aliases = {n for n in names}
        aliases.update({n.split(":")[0] for n in names})
        if self.embedding_model not in aliases:
            raise RuntimeError(
                "Ollama model not found. Run ollama list or set EMBEDDING_MODEL to an existing model."
            )

    async def _create_rag(self, working_dir: str, parser: str) -> RAGAnything:
        if self.embedding_provider != "ollama":
            raise RuntimeError(
                "WebUI defaults to local embeddings only. Set EMBEDDING_PROVIDER=ollama."
            )
        await self._ensure_ollama_model_available()

        embedding_func = EmbeddingFunc(
            embedding_dim=self.embedding_dim,
            max_token_size=8192,
            func=partial(
                ollama_embed_runtime,
                model=self.embedding_model,
                host=self.ollama_host,
                target_dim=self.embedding_dim,
            ),
        )
        config = RAGAnythingConfig(
            working_dir=working_dir,
            parser=parser,
            parse_method="auto",
            enable_image_processing=True,
            enable_table_processing=True,
            enable_equation_processing=True,
        )
        return RAGAnything(
            config=config,
            llm_model_func=partial(
                llm_call_runtime,
                model=self.llm_model,
                api_key=self.api_key,
                base_url=self.base_url,
            ),
            vision_model_func=partial(
                vision_call_runtime,
                vision_model=self.vision_model,
                llm_model=self.llm_model,
                api_key=self.api_key,
                base_url=self.base_url,
            ),
            embedding_func=embedding_func,
            lightrag_kwargs={
                "embedding_func_max_async": 1,
                "embedding_batch_num": 1,
                "llm_model_max_async": 1,
                "max_parallel_insert": 1,
            },
        )

    def _save_index_failure(
        self,
        *,
        existing: Optional[DocumentRecord],
        doc_id: str,
        src: Path,
        stored_path: str,
        workdir: Path,
        parser: str,
        pdf_mode: str,
        sha256: str,
        error_message: str,
    ) -> None:
        record = DocumentRecord(
            doc_id=doc_id,
            original_filename=src.name,
            stored_file_rel=self._as_rel(Path(stored_path), self.paths.storage_root),
            working_dir_rel=self._as_rel(workdir, self.paths.storage_root),
            file_type=src.suffix.lower(),
            parser=parser,
            pdf_mode=pdf_mode,
            indexed_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
            status="failed",
            summary="",
            sha256=sha256,
            source_metadata={
                "embedding_provider": self.embedding_provider,
                "embedding_model": self.embedding_model,
                "embedding_dim": self.embedding_dim,
            },
            visual_targets=(existing.visual_targets if existing else []),
            needs_reprocess=True,
            needs_reprocess_reason=error_message,
            error_message=error_message,
        )
        if existing:
            self.registry = [r for r in self.registry if r.doc_id != existing.doc_id]
        self.registry.append(record)
        self.registry_store.save(self.registry)
        self.rag_cache.pop(doc_id, None)

    async def _process_single_document(
        self,
        file_path: str,
        force_parser: Optional[str] = None,
        pdf_page_range: Optional[str] = None,
        vision_target: Optional[str] = None,
        vision_page_range: Optional[str] = None,
        max_vision_pages: int = 1,
        no_vision: bool = True,
        skip_kg_extraction: bool = True,
        force_reprocess: bool = False,
    ) -> str:
        from examples import raganything_example as rx

        src = Path(file_path)
        if not src.exists():
            return f"Missing file: {file_path}"

        parser = (force_parser or detect_parser_for_file(str(src))).lower()
        sha256 = self._sha256_file(str(src))
        existing = self._find_by_hash(sha256)
        if existing:
            existing_ok, _ = self._record_index_status(existing)
        else:
            existing_ok = False
        if existing and existing_ok and (not force_reprocess):
            return (
                f"Skipped duplicate: {existing.original_filename} "
                f"(already indexed as {existing.doc_id})"
            )

        doc_id = existing.doc_id if existing else uuid.uuid4().hex[:12]
        stored_path = self._persist_upload_file(str(src), doc_id)
        workdir = self.paths.docs_root / doc_id
        if force_reprocess and workdir.exists():
            shutil.rmtree(workdir, ignore_errors=True)
        workdir.mkdir(parents=True, exist_ok=True)

        pdf_mode = "auto"
        if parser == "pdf_hybrid":
            parser = "pdf_fast"
            pdf_mode = "hybrid"
            logger.info(
                "Parser alias pdf_hybrid resolved to pdf_fast with pdf_mode=hybrid"
            )
        if parser == "pdf_fast":
            pdf_mode = "hybrid"

        parser_for_run = (
            "pdf_hybrid" if parser == "pdf_fast" and pdf_mode == "hybrid" else parser
        )
        file_type = src.suffix.lower()
        stored_abs = str(Path(stored_path).resolve())
        working_abs = str(workdir.resolve())
        logger.info(
            "Index request: file=%s type=%s parser=%s pdf_mode=%s stored=%s working_dir=%s",
            str(src.resolve()),
            file_type,
            parser_for_run,
            pdf_mode,
            stored_abs,
            working_abs,
        )
        logger.info(
            "Index options: no_vision=%s vision_target=%s vision_page_range=%s max_vision_pages=%s skip_kg_extraction=%s",
            no_vision,
            vision_target,
            vision_page_range,
            max_vision_pages,
            skip_kg_extraction,
        )
        logger.info("Working dir exists before process: %s", workdir.exists())

        if file_type == ".pptx":
            try:
                import docling  # type: ignore  # noqa: F401
            except Exception:
                err = "PPTX indexing is not available in this Docker build yet."
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=err,
                )
                return f"Indexing failed for {src.name}: {err}"
        if file_type == ".pdf":
            dep_err = self._check_pdf_dependencies()
            if dep_err:
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=dep_err,
                )
                return f"Indexing failed for {src.name}: {dep_err}"

        try:
            if file_type == ".docx" and parser_for_run == "simple_docx":
                try:
                    import docling  # type: ignore  # noqa: F401

                    docling_available = True
                except Exception:
                    docling_available = False
                if docling_available:
                    with self._patched_env_for_ingest():
                        await rx.process_with_rag(
                            file_path=stored_path,
                            output_dir="./output",
                            api_key=self.api_key,
                            base_url=self.base_url,
                            working_dir=working_abs,
                            parser=parser_for_run,
                            parse_method="auto",
                            mineru_backend="pipeline",
                            mineru_device="cpu",
                            fallback_parser="docling",
                            enable_parser_fallback=False,
                            max_chars=None,
                            max_paragraphs=None,
                            skip_query=True,
                            cli_queries=None,
                            embedding_workers=1,
                            embedding_batch_num=1,
                            embedding_max_retries=1,
                            embedding_backoff_base_sec=8.0,
                            embedding_backoff_max_sec=30.0,
                            llm_max_retries=2,
                            llm_backoff_base_sec=6.0,
                            llm_backoff_max_sec=30.0,
                            pdf_mode=pdf_mode,
                            max_pages=None,
                            page_range=pdf_page_range,
                            vision_page_range=vision_page_range,
                            max_vision_pages=max_vision_pages,
                            vision_target=vision_target,
                            no_vision=no_vision,
                            skip_kg_extraction=skip_kg_extraction,
                            raise_on_failure=True,
                        )
                else:
                    logger.info(
                        "DOCX simple_docx/docling unavailable in Docker; using python-docx fallback"
                    )
                    with self._patched_env_for_ingest():
                        await self._index_docx_with_python_docx(
                            file_path=stored_path,
                            workdir=workdir,
                        )
            else:
                with self._patched_env_for_ingest():
                    await rx.process_with_rag(
                        file_path=stored_path,
                        output_dir="./output",
                        api_key=self.api_key,
                        base_url=self.base_url,
                        working_dir=working_abs,
                        parser=parser_for_run,
                        parse_method="auto",
                        mineru_backend="pipeline",
                        mineru_device="cpu",
                        fallback_parser="docling",
                        enable_parser_fallback=False,
                        max_chars=None,
                        max_paragraphs=None,
                        skip_query=True,
                        cli_queries=None,
                        embedding_workers=1,
                        embedding_batch_num=1,
                        embedding_max_retries=1,
                        embedding_backoff_base_sec=8.0,
                        embedding_backoff_max_sec=30.0,
                        llm_max_retries=2,
                        llm_backoff_base_sec=6.0,
                        llm_backoff_max_sec=30.0,
                        pdf_mode=pdf_mode,
                        max_pages=None,
                        page_range=pdf_page_range,
                        vision_page_range=vision_page_range,
                        max_vision_pages=max_vision_pages,
                        vision_target=vision_target,
                        no_vision=no_vision,
                        skip_kg_extraction=skip_kg_extraction,
                        raise_on_failure=True,
                    )
        except Exception as exc:
            if "Content already exists" in str(exc):
                user_msg = (
                    "This file already exists in the current index. "
                    "Please use Reprocess, which will clear the old index first."
                )
                if not force_reprocess:
                    self._save_index_failure(
                        existing=existing,
                        doc_id=doc_id,
                        src=src,
                        stored_path=stored_path,
                        workdir=workdir,
                        parser=parser_for_run,
                        pdf_mode=pdf_mode,
                        sha256=sha256,
                        error_message=user_msg,
                    )
                    return f"Indexing failed for {src.name}: {user_msg}"
            err = (
                f"file_type={file_type}, parser={parser_for_run}, "
                f"working_dir={workdir.resolve()}, error={exc}"
            )
            logger.error("Indexing failed: %s\n%s", err, traceback.format_exc())
            self._save_index_failure(
                existing=existing,
                doc_id=doc_id,
                src=src,
                stored_path=stored_path,
                workdir=workdir,
                parser=parser_for_run,
                pdf_mode=pdf_mode,
                sha256=sha256,
                error_message=err,
            )
            return f"Indexing failed for {src.name}: {exc}"

        artifact_status = index_artifact_status(str(workdir.resolve()))
        logger.info(
            "Working dir exists after process: %s",
            artifact_status["working_dir_exists"],
        )
        logger.info(
            "Working dir contents after process: %s", artifact_status["dir_entries"]
        )
        logger.info("Index artifact status: %s", artifact_status["files"])

        if not self._is_index_ready(str(workdir.resolve())):
            status_err = read_doc_status_error(str(workdir.resolve()), src.name)
            if status_err:
                lowered = status_err.lower()
                if any(
                    k in lowered
                    for k in ["quota", "resource_exhausted", "rate limit", "429"]
                ):
                    err = (
                        "LLM quota exceeded while extracting entities/chunks. "
                        "Try again later or reduce document size via WEBUI_DOCX_MAX_CHARS "
                        "(default 12000 for DOCX fallback). "
                        f"Detail: {status_err}"
                    )
                else:
                    err = f"Document processing failed: {status_err}"
                self._save_index_failure(
                    existing=existing,
                    doc_id=doc_id,
                    src=src,
                    stored_path=stored_path,
                    workdir=workdir,
                    parser=parser_for_run,
                    pdf_mode=pdf_mode,
                    sha256=sha256,
                    error_message=err,
                )
                return f"Indexing failed for {src.name}: {err}"
            err = (
                f"process completed but working_dir index artifacts were not created. "
                f"expected={workdir.resolve()} "
                f"(file_type={file_type}, parser={parser_for_run})"
            )
            self._save_index_failure(
                existing=existing,
                doc_id=doc_id,
                src=src,
                stored_path=stored_path,
                workdir=workdir,
                parser=parser_for_run,
                pdf_mode=pdf_mode,
                sha256=sha256,
                error_message=err,
            )
            return f"Indexing failed for {src.name}: {err}"

        record = DocumentRecord(
            doc_id=doc_id,
            original_filename=src.name,
            stored_file_rel=self._as_rel(Path(stored_path), self.paths.storage_root),
            working_dir_rel=self._as_rel(workdir, self.paths.storage_root),
            file_type=src.suffix.lower(),
            parser="pdf_hybrid" if parser == "pdf_fast" else parser,
            pdf_mode=pdf_mode,
            indexed_at=time.strftime("%Y-%m-%dT%H:%M:%S"),
            status="indexed",
            summary="",
            sha256=sha256,
            source_metadata={
                "embedding_provider": self.embedding_provider,
                "embedding_model": self.embedding_model,
                "embedding_dim": self.embedding_dim,
            },
            visual_targets=(existing.visual_targets if existing else []),
            needs_reprocess=False,
            needs_reprocess_reason="",
            error_message="",
        )

        if existing:
            self.registry = [r for r in self.registry if r.doc_id != existing.doc_id]
        self.registry.append(record)
        self.registry_store.save(self.registry)
        self.rag_cache.pop(doc_id, None)

        note = ""
        if record.file_type in {".png", ".jpg", ".jpeg"} and no_vision:
            note = (
                "\nno_vision applies to PDF visual page rendering only; "
                "direct image input still uses vision."
            )
        return f"Indexed: {src.name} [{doc_id}]\nWorking dir: {workdir}{note}"

    async def process_index(
        self,
        file_paths: list[str] | str,
        force_parser: Optional[str] = None,
        pdf_page_range: Optional[str] = None,
        vision_target: Optional[str] = None,
        vision_page_range: Optional[str] = None,
        max_vision_pages: int = 1,
        no_vision: bool = True,
        skip_kg_extraction: bool = True,
        force_reprocess: bool = False,
    ) -> str:
        if isinstance(file_paths, str):
            file_paths = [file_paths]
        if not file_paths:
            return "Please upload/select files first."
        if not self.api_key:
            return (
                "Missing API key. Set LLM_BINDING_API_KEY "
                "or GEMINI_API_KEY/GOOGLE_API_KEY."
            )

        t0 = time.perf_counter()
        logs = []
        for fp in file_paths:
            try:
                logs.append(
                    await self._process_single_document(
                        file_path=fp,
                        force_parser=force_parser,
                        pdf_page_range=pdf_page_range,
                        vision_target=vision_target,
                        vision_page_range=vision_page_range,
                        max_vision_pages=max_vision_pages,
                        no_vision=no_vision,
                        skip_kg_extraction=skip_kg_extraction,
                        force_reprocess=force_reprocess,
                    )
                )
            except Exception as exc:
                if self._is_gemini_quota_exceeded(exc):
                    logs.append(
                        "Gemini quota exceeded. Continue with text/table-only mode "
                        "or wait for quota reset."
                    )
                else:
                    logs.append(f"Failed indexing {Path(fp).name}: {exc}")

        self.state.processed = any(
            ("Indexed:" in x) or ("Skipped duplicate:" in x) for x in logs
        )
        self.state.last_process_ts = time.time()
        self.state.ingest_count += len(file_paths)
        elapsed = time.perf_counter() - t0
        self.state.processing_summary = (
            "\n\n".join(logs) + f"\n\nTotal processing time: {elapsed:.2f}s"
        )
        return self.state.processing_summary

    async def analyze_visual_target(
        self,
        doc_id: str,
        vision_target: Optional[str],
        vision_page_range: Optional[str],
        max_vision_pages: int,
    ) -> str:
        rec = self._find_by_id(doc_id)
        if rec is None:
            return "Please select an indexed file first."
        if rec.file_type != ".pdf":
            return "Visual target analysis is intended for PDF files."

        result = await self._process_single_document(
            file_path=str(self._resolve_record_stored_file(rec)),
            force_parser="pdf_hybrid",
            pdf_page_range=None,
            vision_target=vision_target,
            vision_page_range=vision_page_range,
            max_vision_pages=max_vision_pages,
            no_vision=False,
            skip_kg_extraction=True,
        )
        rec2 = self._find_by_id(doc_id)
        if rec2 is not None:
            rec2.visual_targets.append(
                {
                    "vision_target": vision_target or "",
                    "vision_page_range": vision_page_range or "",
                    "max_vision_pages": int(max_vision_pages),
                    "processed_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                }
            )
            self.registry_store.save(self.registry)
        return result


class DocumentLifecycleService(DocumentRegistryService, IndexingServiceMixin):
    def __init__(self, paths):
        super().__init__(paths)
        self.state = UIState()

        embedding_cfg = resolve_embedding_runtime_config(default_provider="ollama")
        self.embedding_provider = embedding_cfg.provider
        self.embedding_model = embedding_cfg.model
        self.embedding_dim = embedding_cfg.dim
        self.ollama_host = embedding_cfg.ollama_host
        self.embedding_cfg = embedding_cfg

        self.llm_model = os.getenv("LLM_MODEL", "gemini-3.1-flash-lite")
        self.vision_model = os.getenv("VISION_MODEL", "gemini-3.1-flash-lite")
        self.fallback_vision_model = os.getenv(
            "FALLBACK_VISION_MODEL", "gemini-2.5-flash-lite"
        )
        self.api_key = (
            os.getenv("LLM_BINDING_API_KEY")
            or os.getenv("GEMINI_API_KEY")
            or os.getenv("GOOGLE_API_KEY")
        )
        self.base_url = os.getenv("LLM_BINDING_HOST")

    @staticmethod
    def _is_gemini_quota_exceeded(exc: Exception) -> bool:
        msg = str(exc)
        return (
            "429" in msg
            or "RESOURCE_EXHAUSTED" in msg
            or "quota" in msg.lower()
            or "rate limit" in msg.lower()
        )

    @staticmethod
    def _sha256_file(path: str) -> str:
        return sha256_file(path)

    @staticmethod
    def _is_index_ready(working_dir: str) -> bool:
        return is_index_ready(working_dir)

    def _source_metadata(self) -> dict[str, Any]:
        return {
            "embedding_provider": self.embedding_provider,
            "embedding_model": self.embedding_model,
            "embedding_dim": self.embedding_dim,
        }

    def upload_document_bytes(self, *, filename: str, content: bytes) -> dict[str, Any]:
        if not filename:
            raise RuntimeError("Upload failed: missing filename.")
        safe_name = safe_filename(Path(filename).name)
        file_type = Path(safe_name).suffix.lower()
        sha256 = hashlib.sha256(content).hexdigest()
        existing = self._find_by_hash(sha256)
        if existing is not None:
            ok, _ = self._record_index_status(existing)
            message = (
                "Duplicate document already indexed. Returning existing record."
                if ok
                else "Duplicate document already uploaded. Returning existing record."
            )
            if not ok and existing.status == "indexed":
                existing.needs_reprocess = True
                existing.needs_reprocess_reason = "index artifacts missing"
                self.save_registry()
            return {
                "ok": True,
                "document": self.serialize_document_result(existing),
                "message": message,
            }

        doc_id = uuid.uuid4().hex[:12]
        self.paths.uploads_dir.mkdir(parents=True, exist_ok=True)
        stored_path = self.paths.uploads_dir / f"{doc_id}_{safe_name}"
        stored_path.write_bytes(content)
        working_dir = self.paths.docs_root / doc_id
        record = self.create_uploaded_record(
            doc_id=doc_id,
            filename=safe_name,
            stored_file_rel=self._as_rel(stored_path, self.paths.storage_root),
            working_dir_rel=self._as_rel(working_dir, self.paths.storage_root),
            file_type=file_type,
            sha256=sha256,
            status="uploaded",
            parser=detect_parser_for_file(safe_name),
            pdf_mode="hybrid" if file_type == ".pdf" else "auto",
            source_metadata=self._source_metadata(),
        )
        self.upsert_record(record)
        return {
            "ok": True,
            "document": self.serialize_document_result(record),
            "message": "Document uploaded successfully.",
        }

    async def index_document_by_id(
        self, doc_id: str, *, force_reprocess: bool = False
    ) -> dict[str, Any]:
        rec = self._find_by_id(doc_id)
        if rec is None:
            return {
                "ok": False,
                "document": None,
                "error": "Document not found.",
                "status_code": 404,
            }

        ok, _ = self._record_index_status(rec)
        if ok and not force_reprocess:
            return {
                "ok": True,
                "document": self.serialize_document_result(rec),
                "message": "Document already indexed.",
            }

        if not rec.stored_file_rel:
            rec.status = "failed"
            rec.error_message = "Stored file path is missing."
            rec.needs_reprocess = True
            rec.needs_reprocess_reason = rec.error_message
            self.save_registry()
            return {
                "ok": False,
                "document": self.serialize_document_result(rec),
                "error": f"Indexing failed: {self._sanitize_error_text(rec.error_message)}",
            }

        source_file = self._resolve_record_stored_file(rec)
        if not source_file.exists():
            rec.status = "failed"
            rec.error_message = f"Stored file is missing: {rec.stored_file_rel}"
            rec.needs_reprocess = True
            rec.needs_reprocess_reason = rec.error_message
            self.save_registry()
            return {
                "ok": False,
                "document": self.serialize_document_result(rec),
                "error": f"Indexing failed: {self._sanitize_error_text(rec.error_message)}",
            }

        if force_reprocess:
            workdir = self._resolve_record_working_dir(rec)
            if workdir.exists():
                shutil.rmtree(workdir, ignore_errors=True)
            workdir.mkdir(parents=True, exist_ok=True)
            self.rag_cache.pop(doc_id, None)

        result = await self._process_single_document(
            file_path=str(source_file),
            force_parser=rec.parser or None,
            pdf_page_range=None,
            vision_target=None,
            vision_page_range=None,
            max_vision_pages=1,
            no_vision=True,
            skip_kg_extraction=True,
            force_reprocess=force_reprocess,
        )
        self._reload_registry()
        updated = self._find_by_id(doc_id)
        if updated is None:
            return {
                "ok": False,
                "document": None,
                "error": "Indexing failed: registry record missing after processing.",
            }
        if updated.status == "indexed" and self._record_index_status(updated)[0]:
            return {
                "ok": True,
                "document": self.serialize_document_result(updated),
                "message": (
                    "Document reprocessed successfully."
                    if force_reprocess
                    else "Document indexed successfully."
                ),
            }
        return {
            "ok": False,
            "document": self.serialize_document_result(updated),
            "error": f"Indexing failed: {self._sanitize_error_text(updated.error_message or result)}",
        }
